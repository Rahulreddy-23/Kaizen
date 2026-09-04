import csv

import openpyxl
import pytest

from kaizen.terminology.exchange import export_csv, export_xlsx, import_csv, import_xlsx
from kaizen.terminology.repository import TerminologyRepository


@pytest.fixture
def repo(tmp_path):
    r = TerminologyRepository(tmp_path / "kaizen.db")
    r.initialize()
    return r


def test_xlsx_export_has_readable_columns(repo, tmp_path):
    path = export_xlsx(repo, tmp_path / "rels.xlsx")
    ws = openpyxl.load_workbook(path)["Relationships"]
    header = [c.value for c in ws[1]]
    assert header[:6] == ["ID", "Canonical", "Aliases", "Scope", "Doc Types", "Item Anchors"]
    assert "Version" in header and "Active" in header and "Notes" in header
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert any(r[0] == "REL-001" and "TAPE ANCHOR PER-Q-CATH" in r[2] for r in rows)


def test_xlsx_round_trip_updates_and_creates(repo, tmp_path):
    path = export_xlsx(repo, tmp_path / "rels.xlsx")
    wb = openpyxl.load_workbook(path)
    ws = wb["Relationships"]
    header = [c.value for c in ws[1]]
    col = {h: i + 1 for i, h in enumerate(header)}
    for row in ws.iter_rows(min_row=2):
        if row[col["ID"] - 1].value == "REL-001":
            row[col["Aliases"] - 1].value = "TAPE ANCHOR PER-Q-CATH; TAPE ANCHOR; TAPE, SURGICAL; ANCHOR TAPE"
            row[col["Notes"] - 1].value = "edited in Excel"
    ws.append([None, "Vessel Dilator", "DILATOR VESSEL; DILATOR", "global", "", "", "imported", "quality", "Y", "new from Excel"])
    wb.save(path)
    result = import_xlsx(repo, path, imported_by="quality")
    assert result.updated == 1 and result.created == 1 and result.unchanged >= 5
    rel = repo.get("REL-001")
    assert "ANCHOR TAPE" in rel.aliases and rel.version == 2 and rel.notes == "edited in Excel"
    new = [r for r in repo.list() if r.canonical == "Vessel Dilator"][0]
    assert new.aliases == ["DILATOR VESSEL", "DILATOR"] and new.provenance == "imported" and new.created_by == "quality"


def test_csv_round_trip(repo, tmp_path):
    path = export_csv(repo, tmp_path / "rels.csv")
    with open(path, newline="", encoding="utf-8") as fh:
        rows = list(csv.DictReader(fh))
    assert rows[0]["ID"] == "REL-001"
    with open(path, "a", newline="", encoding="utf-8") as fh:
        w = csv.DictWriter(fh, fieldnames=rows[0].keys())
        w.writerow({"ID": "", "Canonical": "Lubricating Jelly", "Aliases": "JELLY LUBRICATING", "Scope": "global", "Doc Types": "", "Item Anchors": "", "Provenance": "imported", "Created By": "csv", "Active": "Y", "Notes": ""})
    result = import_csv(repo, path, imported_by="csv")
    assert result.created == 1
    assert any(r.canonical == "Lubricating Jelly" for r in repo.list())


def test_import_rejects_bad_scope(repo, tmp_path):
    path = tmp_path / "bad.csv"
    path.write_text("ID,Canonical,Aliases,Scope,Doc Types,Item Anchors,Provenance,Created By,Active,Notes\n,Thing,THING,planet:mars,,,manual,x,Y,\n")
    result = import_csv(repo, path, imported_by="x")
    assert result.created == 0 and result.errors and "scope" in result.errors[0].lower()
