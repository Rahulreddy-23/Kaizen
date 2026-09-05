"""Optional Excel round-trip: a reviewer records decisions in the exported workbook and imports them back.

Rules under test: the workbook must belong to the run; only the importing reviewer's own columns are read;
values are validated; a decision the database changed after the export is a conflict, never a silent
overwrite; a dry run reports everything and writes nothing; every applied decision is audited.
"""

import openpyxl
import pytest

from kaizen.datasets.build import build_golden
from kaizen.models import Thresholds
from kaizen.pipeline import run_folder, save_run
from kaizen.reporting.excel import export_with_review
from kaizen.reporting.excel_import import import_decisions
from kaizen.review.store import ReviewStore
from kaizen.workspace import Workspace

SHEET = "BOM_Label"


@pytest.fixture()
def env(tmp_path):
    golden = build_golden(tmp_path / "golden")
    ws = Workspace(tmp_path / "ws")
    run = run_folder(golden, ws.repository.store(), Thresholds())
    ws.register_run(run, save_run(run, ws.runs_dir / run.metadata.run_id / "run.json"))
    return ws, run


def _export(ws, run, path):
    return export_with_review(ws, run, path)


def _cols(sheet):
    header = [c.value for c in sheet[1]]
    return {name: i + 1 for i, name in enumerate(header)}


def _edit(path, edits: dict[str, dict[str, str]], sheet_name: str = SHEET, extra_rows: list[dict] | None = None):
    """edits: row_id → {column name → value}. extra_rows: appended rows (dict of column → value)."""
    wb = openpyxl.load_workbook(path)
    sh = wb[sheet_name]
    cols = _cols(sh)
    for r in range(2, sh.max_row + 1):
        rid = sh.cell(row=r, column=cols["Row ID"]).value
        if rid in edits:
            for col, val in edits[rid].items():
                sh.cell(row=r, column=cols[col], value=val)
    for extra in extra_rows or []:
        r = sh.max_row + 1
        for col, val in extra.items():
            sh.cell(row=r, column=cols[col], value=val)
    wb.save(path)


def _mismatch_rows(run, n=3):
    return [r.row_id for r in run.results if r.check.value == "BOM_LABEL" and r.role == "item" and r.classification.value == "MISMATCH"][:n]


def test_workbook_for_another_run_is_refused(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    wb = openpyxl.load_workbook(path)
    meta = wb["Run_Metadata"]
    for row in meta.iter_rows(min_row=1, max_row=6):
        if row[0].value == "Run ID":
            row[1].value = "run-deadbeef0000"
    wb.save(path)
    with pytest.raises(ValueError, match="run-deadbeef0000"):
        import_decisions(ws, run, path, slot=1, reviewer="Dharma")


def test_dry_run_reports_what_would_change_and_writes_nothing(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    r1, r2 = _mismatch_rows(run, 2)
    _edit(path, {r1: {"Reviewer Decision": "CONFIRM_DISCREPANCY", "Reviewer Comment": "qty wrong"}, r2: {"Reviewer Decision": "ACCEPT"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma", dry_run=True)
    assert res.dry_run and sorted(res.applied) == sorted([r1, r2]) and res.conflicts == [] and res.invalid == []
    assert ReviewStore(ws.db).all_decisions(run.metadata.run_id) == {}
    assert "dry run" in res.summary().lower()


def test_import_applies_valid_decisions_skips_blanks_and_reports_unknown_rows(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    r1, r2 = _mismatch_rows(run, 2)
    _edit(path, {r1: {"Reviewer Decision": "CONFIRM_DISCREPANCY", "Reviewer Comment": "qty wrong"}, r2: {"Reviewer Decision": "NEEDS_MORE_INFORMATION"}}, extra_rows=[{"Row ID": "R-nope-001", "Reviewer Decision": "ACCEPT"}])
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma Reddy")
    assert sorted(res.applied) == sorted([r1, r2]) and res.unknown_rows == ["R-nope-001"]
    review = ReviewStore(ws.db)
    d = review.decisions(run.metadata.run_id, r1)[1]
    assert d.decision == "CONFIRM_DISCREPANCY" and d.comment == "qty wrong" and d.reviewer == "Dharma Reddy" and d.seconds_spent is None
    assert review.decisions(run.metadata.run_id, r2)[1].decision == "NEEDS_MORE_INFORMATION"
    assert len(review.all_decisions(run.metadata.run_id)) == 2, "blank decision cells are skipped"
    actions = [r["action"] for r in ws.db.conn.execute("SELECT action FROM audit")]
    assert "review.import" in actions and actions.count("review.decision.slot1") == 2


def test_override_round_trips_with_its_classification(env, tmp_path):
    ws, run = env
    review = ReviewStore(ws.db)
    r1 = _mismatch_rows(run, 1)[0]
    review.decide(run.metadata.run_id, r1, 1, "Dharma", "OVERRIDE", override_classification="EQUIVALENT", comment="same part")
    path = _export(ws, run, tmp_path / "r.xlsx")
    sh = openpyxl.load_workbook(path)[SHEET]
    cols = _cols(sh)
    cell = next(sh.cell(row=r, column=cols["Reviewer Decision"]).value for r in range(2, sh.max_row + 1) if sh.cell(row=r, column=cols["Row ID"]).value == r1)
    assert cell == "OVERRIDE→EQUIVALENT", "the override classification must survive the export"
    same = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert same.unchanged == [r1] and same.applied == []
    _edit(path, {r1: {"Reviewer Decision": "OVERRIDE→MISMATCH"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert res.applied == [r1]
    d = review.decisions(run.metadata.run_id, r1)[1]
    assert d.decision == "OVERRIDE" and d.override_classification == "MISMATCH"


def test_invalid_values_are_reported_not_applied(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    r1, r2, r3 = _mismatch_rows(run, 3)
    _edit(path, {r1: {"Reviewer Decision": "MAYBE"}, r2: {"Reviewer Decision": "OVERRIDE"}, r3: {"Reviewer Decision": "OVERRIDE→BOGUS"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert res.applied == [] and sorted(i["row_id"] for i in res.invalid) == sorted([r1, r2, r3])
    assert all(i["reason"] for i in res.invalid)
    assert ReviewStore(ws.db).all_decisions(run.metadata.run_id) == {}


def test_a_decision_changed_in_the_database_after_the_export_is_a_conflict(env, tmp_path):
    ws, run = env
    review = ReviewStore(ws.db)
    r1 = _mismatch_rows(run, 1)[0]
    path = _export(ws, run, tmp_path / "r.xlsx")
    review.decide(run.metadata.run_id, r1, 1, "Dharma", "ACCEPT", comment="decided in the UI after the export", now="2099-01-01T00:00:00+00:00")
    _edit(path, {r1: {"Reviewer Decision": "CONFIRM_DISCREPANCY", "Reviewer Comment": "decided offline"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert res.applied == [] and len(res.conflicts) == 1
    c = res.conflicts[0]
    assert c["row_id"] == r1 and c["workbook"] == "CONFIRM_DISCREPANCY" and c["database"] == "ACCEPT"
    assert review.decisions(run.metadata.run_id, r1)[1].decision == "ACCEPT", "a conflict never overwrites"
    forced = import_decisions(ws, run, path, slot=1, reviewer="Dharma", force=True)
    assert forced.applied == [r1] and review.decisions(run.metadata.run_id, r1)[1].decision == "CONFIRM_DISCREPANCY"


def test_a_database_decision_older_than_the_export_is_simply_updated(env, tmp_path):
    ws, run = env
    review = ReviewStore(ws.db)
    r1 = _mismatch_rows(run, 1)[0]
    review.decide(run.metadata.run_id, r1, 1, "Dharma", "ACCEPT", now="2000-01-01T00:00:00+00:00")
    path = _export(ws, run, tmp_path / "r.xlsx")
    _edit(path, {r1: {"Reviewer Decision": "CONFIRM_DISCREPANCY"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert res.applied == [r1] and res.conflicts == []


def test_only_the_importing_reviewers_columns_are_read(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    r1 = _mismatch_rows(run, 1)[0]
    _edit(path, {r1: {"Reviewer 2 Decision": "ACCEPT", "Reviewer 2 Comment": "not mine to write"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    assert res.applied == [] and ReviewStore(ws.db).all_decisions(run.metadata.run_id) == {}
    res2 = import_decisions(ws, run, path, slot=2, reviewer="Hemant")
    assert res2.applied == [r1]
    d = ReviewStore(ws.db).decisions(run.metadata.run_id, r1)[2]
    assert d.reviewer == "Hemant" and d.blind is False


def test_summary_reads_as_a_sentence(env, tmp_path):
    ws, run = env
    path = _export(ws, run, tmp_path / "r.xlsx")
    r1 = _mismatch_rows(run, 1)[0]
    _edit(path, {r1: {"Reviewer Decision": "ACCEPT"}})
    res = import_decisions(ws, run, path, slot=1, reviewer="Dharma")
    text = res.summary()
    assert "1 applied" in text and "slot 1" in text and "Dharma" in text
