import openpyxl
import pytest

from kaizen.datasets.pdf_bom import BomRowSpec, BomSpec, render_bom_pdf
from kaizen.datasets.pdf_label import LabelSpec, render_label_pdf
from kaizen.models import Thresholds
from kaizen.pipeline import run_folder
from kaizen.reporting.excel import BOM_LABEL_COLUMNS, write_report
from kaizen.terminology.store import RelationshipStore

REQUIRED = [
    "Row ID", "SKU", "Check", "Source A File", "Source A Page", "Source A Item", "Source A Description", "Source A Quantity",
    "Source B File", "Source B Page", "Source B Item", "Source B Description", "Source B Quantity", "Normalized A", "Normalized B",
    "Classification", "Score", "Relationship ID", "Discrepancy Type", "Discrepancy Detail", "Severity", "Requires Validation",
    "Reviewer Decision", "Reviewer Comment", "Final Status",
]


@pytest.fixture
def run(tmp_path):
    render_bom_pdf(BomSpec(parent_item="1295108NS", parent_description="KIT A", rows=[BomRowSpec(item="0396447", description="ABSORBENT TOWEL"), BomRowSpec(item="5167473", description="TAPE ANCHOR PER-Q-CATH", qty_per="2.0000"), BomRowSpec(item="8880001", description="LONELY PART")]), tmp_path / "sku-001" / "bom.pdf")
    render_label_pdf(LabelSpec(ref="1295108", product_name="Kit A", contents=["1 Each - Towel, Absorbent", "1 Each - Surgical Tape"]), tmp_path / "sku-001" / "label.pdf")
    return run_folder(tmp_path, RelationshipStore.default(), Thresholds())


def test_required_columns_present_in_order_subset():
    assert [c for c in BOM_LABEL_COLUMNS if c in REQUIRED] == REQUIRED


def test_workbook_sheets_and_rows(run, tmp_path):
    path = write_report(run, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(path)
    assert wb.sheetnames[:13] == ["Summary", "BOM_Label", "BOM_Drawing", "Label_Drawing", "PCO_BOM", "Label_Revision", "Action_Items", "Coverage", "Manual_Checklist", "Documents", "Relationships_Used", "Run_Metadata", "Audit_Log"]
    ws = wb["BOM_Label"]
    header = [c.value for c in ws[1]]
    assert header == BOM_LABEL_COLUMNS
    rows = list(ws.iter_rows(min_row=2, values_only=True))
    assert len(rows) == len(run.results)
    col = {name: i for i, name in enumerate(header)}
    tape = next(r for r in rows if r[col["Source A Item"]] == "5167473")
    assert tape[col["Classification"]] == "MISMATCH"
    assert tape[col["Relationship ID"]] == "REL-001"
    assert tape[col["Discrepancy Type"]] == "QTY_MISMATCH"
    assert tape[col["Severity"]] == "MAJOR"
    assert tape[col["Requires Validation"]] == "Y"
    assert tape[col["Source A Page"]] == 1 and tape[col["Source B Page"]] == 1
    assert tape[col["Final Status"]] == "OPEN"
    assert tape[col["Row ID"]].startswith("R-1295108NS-")
    assert "REL-001" in tape[col["Explanation"]]
    lonely = next(r for r in rows if r[col["Source A Item"]] == "8880001")
    assert lonely[col["Classification"]] == "MISSING" and lonely[col["Source B File"]] in (None, "")


def test_reviewer_columns_have_validation_and_rows_are_styled(run, tmp_path):
    path = write_report(run, tmp_path / "report.xlsx")
    ws = openpyxl.load_workbook(path)["BOM_Label"]
    assert ws.freeze_panes == "A2"
    assert ws.auto_filter.ref
    assert any("ACCEPT" in str(dv.formula1) for dv in ws.data_validations.dataValidation)
    assert len(ws.conditional_formatting) > 0


def test_metadata_sheet_answers_where_did_this_come_from(run, tmp_path):
    path = write_report(run, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(path)
    text = " ".join(str(c.value) for row in wb["Run_Metadata"].iter_rows() for c in row if c.value is not None)
    for f in run.metadata.inputs:
        assert f.sha256 in text
    assert run.metadata.terminology_version in text
    assert "potential" in text and "0.85" in text
    assert "NOT IMPLEMENTED" in text
    rel_text = " ".join(str(c.value) for row in wb["Relationships_Used"].iter_rows() for c in row if c.value is not None)
    assert "REL-001" in rel_text and "Surgical Tape" in rel_text
    docs_text = " ".join(str(c.value) for row in wb["Documents"].iter_rows() for c in row if c.value is not None)
    assert "bom_pdf" in docs_text and run.documents[0].sha256 in docs_text
    assert wb["Audit_Log"].max_row >= 2


def test_summary_counts(run, tmp_path):
    path = write_report(run, tmp_path / "report.xlsx")
    ws = openpyxl.load_workbook(path)["Summary"]
    values = [[c for c in row] for row in ws.iter_rows(values_only=True)]
    flat = [str(v) for row in values for v in row if v is not None]
    assert "1295108NS" in flat
    assert "NEEDS REVIEW" in flat


def test_accuracy_sheet_when_metrics_given(run, tmp_path):
    path = write_report(run, tmp_path / "report.xlsx", metrics={"overall": {"precision": 1.0, "recall": 0.5}, "per_type": {"QTY_MISMATCH": {"tp": 1, "fp": 0, "fn": 1}}})
    wb = openpyxl.load_workbook(path)
    assert "Accuracy" in wb.sheetnames
    text = " ".join(str(c.value) for row in wb["Accuracy"].iter_rows() for c in row if c.value is not None)
    assert "QTY_MISMATCH" in text and "0.5" in text
