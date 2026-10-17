"""Export the saved projection without changing status or losing source entries."""

from openpyxl import load_workbook

from kaizen.reporting.quick_view import export_quick_view
from tests.unit.test_quick_view import fixture, project


def test_comparison_workbook_preserves_status_sources_and_literal_text(tmp_path):
    run = fixture()
    run.documents[0].items[0].description = '=HYPERLINK("https://example.com")'
    run.documents[0].items[0].evidence.raw_text = "=" + "Original\x00 raw " * 4000
    data = project(run)
    path = export_quick_view([data], tmp_path / "comparison.xlsx", run_id="quick", run_name="Synthetic")
    book = load_workbook(path, data_only=False)
    assert book.sheetnames == ["Read me", "Sources", "1 SKU-A"]
    sheet = book["1 SKU-A"]
    assert sheet["A2"].value == "Matched"
    assert sheet["B2"].font.color.rgb == "00216335"
    assert "BOM.pdf · p1" in sheet["B2"].value
    assert sheet.freeze_panes == "B2" and sheet.auto_filter.ref == "A1:E2"
    sources = list(book["Sources"].values)[1:]
    bom = [r for r in sources if r[2] == "BOM"]
    assert len(bom) == 2
    assert "".join(r[11] for r in bom) == run.documents[0].items[0].evidence.raw_text.replace("\x00", "")
    assert bom[0][6] == 1 and bom[0][9] == 1
    assert book["Sources"]["F2"].data_type == "s"
    assert all(c.data_type != "f" for s in book for row in s for c in row)
    assert all("/Users/" not in str(c.value) for s in book for row in s for c in row)


def test_missing_counterpart_is_written_and_coloured_without_false_match(tmp_path):
    run = fixture()
    run.documents = [d for d in run.documents if d.doc_type.value != "LABEL"]
    run.results = [r for r in run.results if r.check == "BOM_DRAWING"]
    data = project(run)
    book = load_workbook(export_quick_view([data], tmp_path / "missing.xlsx", run_id="quick", run_name=""))
    sheet = book["1 SKU-A"]
    assert sheet["A2"].value == "Missing or needs review"
    assert sheet["C2"].value == "Document unavailable"
    assert sheet["C2"].font.color.rgb == "008A4B08"
    assert sheet["B2"].value.startswith("Matched")
    assert len(list(book["Sources"].values)) == 3


def test_all_skus_and_entries_are_exported_without_a_screen_page_limit(tmp_path):
    first = project(fixture())
    first['rows'] = [first['rows'][0] | {'id': f'row-{n}'} for n in range(121)]
    second = project(fixture()) | {'sku': 'SKU/B:second'}
    book = load_workbook(export_quick_view([first, second], tmp_path / 'all.xlsx', run_id='quick', run_name=''))
    assert book['1 SKU-A'].max_row == 122
    assert book['2 SKU_B_second'].max_row == 2
    assert book['Sources'].max_row == 1 + 122 * 3
    assert book['Sources']['A365'].value == 'SKU/B:second'
