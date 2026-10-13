"""SKU downloads stay compact, complete, traceable and compatible with offline review."""

import copy
import json
from pathlib import Path

import openpyxl
import pytest

from kaizen.models import Classification, DocType, Thresholds
from kaizen.pipeline import Ingested, run_checks, run_folder, save_run
from kaizen.reporting.excel import export_with_review, write_report
from kaizen.reporting.excel_import import import_decisions, read_metadata
from kaizen.review.progress import review_snapshot
from kaizen.review.store import ReviewStore
from kaizen.terminology.store import RelationshipStore
from kaizen.workspace import Workspace


@pytest.fixture(scope="module")
def run():
    return run_folder(Path(__file__).resolve().parents[2] / "datasets/golden", RelationshipStore.default(), Thresholds())


def table(sheet):
    header = next(row for row in sheet if any(c.value == "Row ID" for c in row))
    columns = {c.value: c.column for c in header if c.value is not None}
    for display, field in {"Result": "Current Classification", "Status": "Review State", "Finding": "Finding / Next Step", "Item": "Source A Item", "Description A": "Source A Description", "Description B": "Source B Description", "Qty A": "Source A Quantity", "Qty B": "Source B Quantity"}.items():
        if display in columns:
            columns[field] = columns[display]
    records = {str(sheet.cell(row, columns["Row ID"]).value): row for row in range(header[0].row + 1, sheet.max_row + 1) if sheet.cell(row, columns["Row ID"]).value}
    return columns, records


def workspace(tmp_path, run):
    ws = Workspace(tmp_path / "workspace")
    ws.register_run(run, save_run(run, ws.runs_dir / run.metadata.run_id / "run.json"))
    return ws


def test_summary_has_one_visible_tab_per_sku_and_only_attention_rows(run, tmp_path):
    wb = openpyxl.load_workbook(write_report(run, tmp_path / "summary.xlsx", detail="summary"))
    skus = {g.sku for g in run.groups} | {r.sku for r in run.results}
    assert {s.title for s in wb if s.sheet_state == "visible"} == skus
    expected = {r.row_id for r in run.results if r.role != "exempt" and (r.requires_validation or any(d.severity.value != "INFO" for d in r.discrepancies))}
    actual = set()
    for sku in skus:
        sh = wb[sku]
        cols, records = table(sh)
        actual.update(records)
        assert {r.row_id for r in run.results if r.sku == sku} >= records.keys()
        assert "Score" not in cols and "Source A File" not in cols
        assert not any("composite_fuzzy" in str(cell.value) or "token_sort=" in str(cell.value) for row in sh for cell in row)
        assert sh.auto_filter.ref and sh.freeze_panes and sh.sheet_view.showGridLines
        assert sh.column_dimensions[openpyxl.utils.get_column_letter(cols["Row ID"])].hidden
        assert sh.data_validations.dataValidation
        total = next(r for r in sh.iter_rows(max_col=6) if r[0].value == "Total")
        assert total[1].value == sum(r.sku == sku and r.role != "exempt" for r in run.results)
    assert actual == expected
    assert read_metadata(wb)["Run ID"] == run.metadata.run_id
    assert read_metadata(wb)["Export layout"] == "sku"
    assert wb["Run_Metadata"].sheet_state == "hidden"


def test_complete_preserves_every_row_and_source_fields_once(run, tmp_path):
    wb = openpyxl.load_workbook(write_report(run, tmp_path / "complete.xlsx", detail="complete"))
    found = {}
    for sku in {r.sku for r in run.results}:
        sh = wb[sku]
        cols, records = table(sh)
        for rid, index in records.items():
            assert rid not in found
            found[rid] = (sh, cols, index)
    assert found.keys() == {r.row_id for r in run.results}
    for row in run.results:
        sh, cols, index = found[row.row_id]
        value = {name: sh.cell(index, col).value for name, col in cols.items()}.get
        assert value("Classification") == row.classification.value
        assert value("Score") == row.score
        if row.source_a:
            assert value("Source A Description") == row.source_a.description
            assert value("Source A Quantity") == (float(row.source_a.quantity) if row.source_a.quantity is not None else None)
            assert value("Source A Raw Text") == (row.source_a.evidence.raw_text or None)
            assert json.loads(value("Source A Attributes")) == json.loads(json.dumps(row.source_a.attributes, default=str))
    assert {"Documents", "Audit_Log", "Relationships_Used", "Manual_Checklist", "Coverage", "Action_Items"} <= set(wb.sheetnames)
    assert not {"BOM_Label", "BOM_Drawing", "Label_Drawing", "PCO_BOM", "Label_Revision"} & set(wb.sheetnames)


def test_summary_uses_current_progress_without_erasing_engine_results(run, tmp_path):
    ws = workspace(tmp_path, run)
    review = ReviewStore(ws.db)
    cleared, confirmed, pending = [r for r in run.results if r.role == "item" and r.requires_validation][:3]
    for row, decision, override in ((cleared, "OVERRIDE", "EXACT"), (confirmed, "CONFIRM_DISCREPANCY", None)):
        for slot in (1, 2):
            review.decide(run.metadata.run_id, row.row_id, slot, f"Reviewer {slot}", decision, override_classification=override)
        review.finalize(run.metadata.run_id, row.row_id, decision, by="Approver")
    review.decide(run.metadata.run_id, pending.row_id, 1, "Reviewer", "ACCEPT")
    snapshots = {r["row_id"]: r for r in review_snapshot(run, review)[0]}
    wb = openpyxl.load_workbook(export_with_review(ws, run, tmp_path / "summary.xlsx", detail="summary"))
    expected = {r.row_id for r in run.results if r.role != "exempt" and (snapshots[r.row_id]["pending_review"] or snapshots[r.row_id]["unresolved"])}
    actual = {rid for sh in wb if sh.sheet_state == "visible" for rid in table(sh)[1]}
    assert actual == expected and cleared.row_id not in actual and confirmed.row_id in actual and pending.row_id in actual
    full = openpyxl.load_workbook(export_with_review(ws, run, tmp_path / "complete.xlsx", detail="complete"))
    sh = full[cleared.sku]
    cols, records = table(sh)
    assert sh.cell(records[cleared.row_id], cols["Classification"]).value == cleared.classification.value
    assert sh.cell(records[cleared.row_id], cols["Current Classification"]).value == "EXACT"
    assert sh.cell(records[cleared.row_id], cols["Approved By"]).value == "Approver"


@pytest.mark.parametrize("detail", ["summary", "complete"])
def test_decisions_round_trip_from_renamed_sku_tabs(run, tmp_path, detail):
    ws = workspace(tmp_path, run)
    row = next(r for r in run.results if r.classification is Classification.MISMATCH and r.requires_validation)
    path = export_with_review(ws, run, tmp_path / "review.xlsx", detail=detail)
    wb = openpyxl.load_workbook(path)
    sh = wb[row.sku]
    cols, records = table(sh)
    sh.cell(records[row.row_id], cols["Reviewer Decision"], "CONFIRM_DISCREPANCY")
    sh.cell(records[row.row_id], cols["Reviewer Comment"], "Needs correction")
    sh.title = "Renamed by reviewer"
    wb.save(path)
    preview = import_decisions(ws, run, path, 1, "Reviewer", dry_run=True)
    assert preview.applied == [row.row_id]
    assert ReviewStore(ws.db).decisions(run.metadata.run_id, row.row_id) == {}
    result = import_decisions(ws, run, path, 1, "Reviewer")
    assert result.applied == [row.row_id]
    assert ReviewStore(ws.db).decisions(run.metadata.run_id, row.row_id)[1].comment == "Needs correction"


def test_sheet_names_handle_collisions_reserved_names_and_leading_zeros(run, tmp_path):
    changed = copy.deepcopy(run)
    changed.groups = []
    names = ["a/b", "a:b", "A-B", "Run_Metadata", "000123", "x" * 40, "x" * 39 + "y", "'quoted'", "[bad]?*"]
    changed.results = [r.model_copy(update={"sku": name, "row_id": f"edge-{index}"}) for index, (r, name) in enumerate(zip(run.results, names))]
    wb = openpyxl.load_workbook(write_report(changed, tmp_path / "names.xlsx", detail="complete"))
    sku_sheets = wb.worksheets[:len(names)]
    assert len({sh.title.casefold() for sh in sku_sheets}) == len(names)
    assert all(len(sh.title) <= 31 and not any(c in sh.title for c in "[]:*?/\\") for sh in sku_sheets)
    assert wb["000123"].title == "000123"
    assert wb["Run_Metadata"].sheet_state == "hidden"
    assert any(sh.title == "Run_Metadata (2)" for sh in sku_sheets)


@pytest.mark.parametrize("detail", ["summary", "complete"])
def test_download_layout_is_a_plain_compact_table(run, tmp_path, detail):
    wb = openpyxl.load_workbook(write_report(run, tmp_path / "plain.xlsx", detail=detail))
    sheet = wb["1295108FNS"]
    columns, records = table(sheet)
    header = next(row for row in sheet if any(c.value == "Row ID" for c in row))
    visible = [c.value for c in header if not sheet.column_dimensions[c.column_letter].hidden]
    expected = ["Check", "Component", "Result", "Status", "Finding", "Reviewer Decision", "Reviewer Comment"] if detail == "summary" else ["Check", "Item", "Description A", "Description B", "Qty A", "Qty B", "Result", "Status", "Reviewer Decision", "Reviewer Comment"]
    assert visible == expected
    assert header[0].row < 12 and (detail != "complete" or header[0].row == 1)
    assert sheet.sheet_view.zoomScale == 100 and sheet.sheet_properties.tabColor is None
    for index in records.values():
        for cell in sheet[index]:
            assert cell.fill.patternType is None
            assert cell.font.sz == 11
    assert not any("Review this comparison" in str(c.value) or "Confirm whether these descriptions" in str(c.value) or "Engine classifications are" in str(c.value) for row in sheet for c in row)


def test_empty_export_and_group_without_comparisons_are_explicit(run, tmp_path):
    changed = run.model_copy(update={"results": [], "groups": []})
    wb = openpyxl.load_workbook(write_report(changed, tmp_path / "empty.xlsx", detail="summary"))
    assert wb["No SKUs"].sheet_state == "visible"
    assert any(c.value == "No comparisons available for this SKU." for row in wb["No SKUs"] for c in row)


@pytest.mark.parametrize("remaining, missing", [(DocType.BOM, "MISSING_LABEL"), (DocType.LABEL, "MISSING_BOM")])
def test_summary_retains_missing_document_coverage_when_comparisons_are_skipped(run, tmp_path, remaining, missing):
    document = next(d for d in run.documents if d.doc_type is remaining)
    incomplete = run_checks(Ingested(root=Path(run.metadata.input_root), documents=[document]), RelationshipStore.default())
    assert not incomplete.results
    finding = next(c for c in incomplete.coverage if c.status == missing)
    ws = workspace(tmp_path, incomplete)
    path = export_with_review(ws, incomplete, tmp_path / "summary.xlsx", detail="summary")
    wb = openpyxl.load_workbook(path)
    sheet = wb[incomplete.groups[0].sku]
    contents = [cell.value for row in sheet for cell in row]
    assert missing.replace("_", " ").capitalize() in contents
    assert finding.detail in contents
    assert "No comparisons available for this SKU." in contents
    assert not table(sheet)[1]
    assert import_decisions(ws, incomplete, path, 1, "Reviewer", dry_run=True).applied == []


def test_a_decision_changed_while_generating_the_download_remains_a_conflict(run, tmp_path, monkeypatch):
    from kaizen.reporting import excel

    ws = workspace(tmp_path, run)
    review = ReviewStore(ws.db)
    row = next(r for r in run.results if r.requires_validation)
    review.decide(run.metadata.run_id, row.row_id, 1, "Reviewer", "CONFIRM_DISCREPANCY")
    original_write = excel.write_report

    def change_while_writing(*args, **kwargs):
        review.decide(run.metadata.run_id, row.row_id, 1, "Reviewer", "NEEDS_MORE_INFORMATION")
        return original_write(*args, **kwargs)

    monkeypatch.setattr(excel, "write_report", change_while_writing)
    path = export_with_review(ws, run, tmp_path / "review.xlsx", detail="summary")
    result = import_decisions(ws, run, path, 1, "Reviewer")
    assert [c["row_id"] for c in result.conflicts] == [row.row_id]
    assert review.decisions(run.metadata.run_id, row.row_id)[1].decision == "NEEDS_MORE_INFORMATION"


def test_source_formula_text_is_literal_in_both_downloads(run, tmp_path):
    changed = copy.deepcopy(run)
    row = next(r for r in changed.results if r.source_a and r.requires_validation)
    row.source_a.description = '=HYPERLINK("https://example.invalid", "test")'
    row.source_a.item_number = None
    for detail in ("summary", "complete"):
        wb = openpyxl.load_workbook(write_report(changed, tmp_path / f"{detail}.xlsx", detail=detail))
        sh = wb[row.sku]
        cols, records = table(sh)
        name = "Component" if detail == "summary" else "Source A Description"
        cell = sh.cell(records[row.row_id], cols[name])
        assert cell.data_type == "s" and cell.value.startswith("=")
