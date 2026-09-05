"""Excel export carries reviewer decisions, action items and the manual checklist; engine columns stay intact."""

import openpyxl
import pytest

from kaizen.datasets.build import build_golden
from kaizen.models import Classification, Thresholds
from kaizen.pipeline import run_folder, save_run
from kaizen.reporting.excel import BOM_LABEL_COLUMNS, export_with_review
from kaizen.review.action_items import ActionItemStore
from kaizen.review.store import ReviewStore
from kaizen.workspace import Workspace


@pytest.fixture
def ws_run(tmp_path):
    golden = build_golden(tmp_path / "golden")
    ws = Workspace(tmp_path / "ws")
    run = run_folder(golden / "sku-002", ws.repository.store(), Thresholds())
    ws.register_run(run, save_run(run, ws.runs_dir / run.metadata.run_id / "run.json"))
    return ws, run


def test_decisions_action_items_and_checklist_in_workbook(ws_run, tmp_path):
    ws, run = ws_run
    rid = run.metadata.run_id
    review, items = ReviewStore(ws.db), ActionItemStore(ws.db)
    row = next(r for r in run.results if r.classification is Classification.MISMATCH and r.check.value == "BOM_LABEL")
    review.decide(rid, row.row_id, slot=1, reviewer="Dharma", decision="CONFIRM_DISCREPANCY", comment="BOM qty wrong")
    review.decide(rid, row.row_id, slot=2, reviewer="Hemant", decision="CONFIRM_DISCREPANCY", comment="agree", blind=True)
    review.finalize(rid, row.row_id, "CONFIRM_DISCREPANCY", by="Dharma")
    ai = items.create_from_result(run, row, reviewer="Dharma", owner="R&D")
    path = export_with_review(ws, run, tmp_path / "report.xlsx")
    wb = openpyxl.load_workbook(path)
    assert "Action_Items" in wb.sheetnames and "Manual_Checklist" in wb.sheetnames
    sheet = wb["BOM_Label"]
    header = [c.value for c in sheet[1]]
    col = {h: i for i, h in enumerate(header)}
    r = next(x for x in sheet.iter_rows(min_row=2, values_only=True) if x[col["Row ID"]] == row.row_id)
    assert r[col["Classification"]] == "MISMATCH"  # engine recommendation untouched
    assert r[col["Reviewer Decision"]] == "CONFIRM_DISCREPANCY" and r[col["Reviewer Name"]] == "Dharma" and r[col["Reviewer Comment"]] == "BOM qty wrong" and r[col["Reviewer Date"]]
    assert r[col["Reviewer 2 Decision"]] == "CONFIRM_DISCREPANCY" and r[col["Reviewer 2 Name"]] == "Hemant"
    assert r[col["Final Status"]] == "FINALIZED: CONFIRM_DISCREPANCY"
    assert "Review State" in header and r[col["Review State"]] == "FINALIZED"
    assert "Action Item" in header and r[col["Action Item"]] == ai.id
    ai_rows = list(wb["Action_Items"].iter_rows(min_row=2, values_only=True))
    assert ai_rows and ai_rows[0][0] == ai.id and row.row_id in ai_rows[0]
    checklist = " ".join(str(c.value) for r_ in wb["Manual_Checklist"].iter_rows() for c in r_ if c.value)
    assert "dot sticker" in checklist.lower()
    summary = " ".join(str(c.value) for r_ in wb["Summary"].iter_rows() for c in r_ if c.value)
    assert "FINALIZED" in summary and "Estimated" in summary


def test_undecided_rows_keep_open_status(ws_run, tmp_path):
    ws, run = ws_run
    path = export_with_review(ws, run, tmp_path / "report.xlsx")
    sheet = openpyxl.load_workbook(path)["BOM_Label"]
    header = [c.value for c in sheet[1]]
    col = {h: i for i, h in enumerate(header)}
    r = next(sheet.iter_rows(min_row=2, values_only=True))
    assert r[col["Final Status"]] == "OPEN" and r[col["Review State"]] == "ENGINE_RECOMMENDED"
    assert [c for c in BOM_LABEL_COLUMNS if c in header] == BOM_LABEL_COLUMNS


def test_accuracy_sheet_survives_a_later_export_from_the_workspace(tmp_path):
    """`kaizen demo` writes accuracy.json beside the run; a UI export without metrics must pick it up, so the
    Accuracy sheet does not vanish the moment someone exports the workbook again."""
    import json

    from kaizen.datasets.build import build_golden
    from kaizen.evaluation.ground_truth import load_ground_truth
    from kaizen.evaluation.harness import evaluate
    from kaizen.models import Thresholds
    from kaizen.pipeline import run_folder, save_run
    from kaizen.reporting.excel import export_with_review
    from kaizen.workspace import Workspace

    golden = build_golden(tmp_path / "golden")
    ws = Workspace(tmp_path / "ws")
    run = run_folder(golden, ws.repository.store(), Thresholds())
    out = ws.runs_dir / run.metadata.run_id
    ws.register_run(run, save_run(run, out / "run.json"))
    metrics = evaluate(run, load_ground_truth(golden / "ground-truth.json")).to_dict()
    (out / "accuracy.json").write_text(json.dumps(metrics))
    wb = openpyxl.load_workbook(export_with_review(ws, run, out / "report.xlsx"))
    assert "Accuracy" in wb.sheetnames
    wb2 = openpyxl.load_workbook(export_with_review(ws, run, tmp_path / "elsewhere.xlsx"))
    assert "Accuracy" not in wb2.sheetnames, "only a sidecar next to the target workbook is used"
