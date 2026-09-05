import json

import openpyxl
import pytest
from typer.testing import CliRunner

from kaizen.cli.main import app
from kaizen.datasets.build import build_golden

runner = CliRunner()


@pytest.fixture(scope="module")
def golden(tmp_path_factory):
    return build_golden(tmp_path_factory.mktemp("data") / "golden")


def test_help_lists_commands():
    result = runner.invoke(app, ["--help"])
    assert result.exit_code == 0
    for cmd in ("run", "ingest", "check", "report", "eval", "dataset"):
        assert cmd in result.output


def test_run_produces_run_json_and_report(golden, tmp_path):
    out = tmp_path / "out"
    result = runner.invoke(app, ["run", str(golden), "--out", str(out)])
    assert result.exit_code == 0, result.output
    assert (out / "run.json").exists()
    assert (out / "report.xlsx").exists()
    run = json.loads((out / "run.json").read_text())
    n_files = sum(1 for p in golden.rglob("*") if p.suffix in (".pdf", ".xlsx", ".csv"))
    assert len(run["documents"]) == n_files  # every BOM, label (incl. old revisions), drawing and PCO parsed
    assert len(run["groups"]) == 10
    assert run["results"]
    wb = openpyxl.load_workbook(out / "report.xlsx")
    assert "BOM_Label" in wb.sheetnames
    assert "1295108NS" in result.output


def test_ingest_lists_documents(golden):
    result = runner.invoke(app, ["ingest", str(golden / "sku-001")])
    assert result.exit_code == 0, result.output
    assert "BOM" in result.output and "LABEL" in result.output
    assert "1295108NS" in result.output


def test_check_then_report_split_flow(golden, tmp_path):
    out = tmp_path / "split"
    r1 = runner.invoke(app, ["check", str(golden / "sku-002"), "--out", str(out)])
    assert r1.exit_code == 0, r1.output
    assert (out / "run.json").exists() and not (out / "report.xlsx").exists()
    r2 = runner.invoke(app, ["report", str(out / "run.json")])
    assert r2.exit_code == 0, r2.output
    assert (out / "report.xlsx").exists()


def test_eval_writes_metrics_and_prints_precision(golden, tmp_path):
    out = tmp_path / "eval"
    result = runner.invoke(app, ["eval", str(golden), "--out", str(out)])
    assert result.exit_code == 0, result.output
    metrics = json.loads((out / "metrics.json").read_text())
    assert "overall" in metrics and "precision" in metrics["overall"]
    assert "OVERALL" in result.output
    assert (out / "report.xlsx").exists()
    assert "Accuracy" in openpyxl.load_workbook(out / "report.xlsx").sheetnames


def test_dataset_build_command(tmp_path):
    result = runner.invoke(app, ["dataset", "build", "--out", str(tmp_path / "g")])
    assert result.exit_code == 0, result.output
    assert (tmp_path / "g" / "ground-truth.json").exists()


def test_custom_relationships_file_and_threshold(golden, tmp_path):
    rels = tmp_path / "rels.json"
    rels.write_text(json.dumps({"version": 1, "relationships": []}))
    out = tmp_path / "out2"
    result = runner.invoke(app, ["run", str(golden / "sku-001"), "--out", str(out), "--relationships", str(rels), "--potential", "0.9"])
    assert result.exit_code == 0, result.output
    run = json.loads((out / "run.json").read_text())
    assert run["metadata"]["thresholds"]["potential"] == 0.9
    assert run["metadata"]["terminology_count"] == 0


def test_demo_command_runs_golden_and_reports_measured_accuracy(golden, tmp_path):
    result = runner.invoke(app, ["--workspace", str(tmp_path / "ws"), "demo", "--dataset", str(golden), "--out", str(tmp_path / "demo")])
    assert result.exit_code == 0, result.output
    assert "Measured against ground truth: precision 1.000, recall 1.000" in result.output
    assert (tmp_path / "demo" / "report.xlsx").exists()


def test_serve_help_mentions_local_binding():
    result = runner.invoke(app, ["serve", "--help"])
    assert result.exit_code == 0 and "127.0.0.1" in result.output


def test_certificate_diff_and_excel_import_commands(golden, tmp_path):
    out = tmp_path / "out"
    assert runner.invoke(app, ["run", str(golden), "--out", str(out)]).exit_code == 0
    run_json = out / "run.json"

    cert = runner.invoke(app, ["certificate", str(run_json), "--out", str(tmp_path / "cert.pdf")])
    assert cert.exit_code == 0, cert.output
    assert (tmp_path / "cert.pdf").read_bytes().startswith(b"%PDF")
    one = runner.invoke(app, ["certificate", str(run_json), "--sku", "1295108NS", "--out", str(tmp_path / "one.pdf")])
    assert one.exit_code == 0 and (tmp_path / "one.pdf").exists()
    assert runner.invoke(app, ["certificate", str(run_json), "--sku", "nope"]).exit_code != 0

    d = runner.invoke(app, ["diff", str(run_json), str(run_json)])
    assert d.exit_code == 0 and "resolved 0" in d.output and "unchanged" in d.output
    dj = runner.invoke(app, ["diff", str(run_json), str(run_json), "--json"])
    assert json.loads(dj.output)["counts"]["new"] == 0

    wb = openpyxl.load_workbook(out / "report.xlsx")
    sh = wb["BOM_Label"]
    cols = {c.value: i + 1 for i, c in enumerate(sh[1])}
    sh.cell(row=2, column=cols["Reviewer Decision"], value="ACCEPT")
    wb.save(out / "report.xlsx")
    dry = runner.invoke(app, ["review", "import", str(run_json), str(out / "report.xlsx"), "--slot", "1", "--reviewer", "Dharma", "--dry-run"])
    assert dry.exit_code == 0 and "1 applied" in dry.output and "Dry run" in dry.output, dry.output
    real = runner.invoke(app, ["review", "import", str(run_json), str(out / "report.xlsx"), "--slot", "1", "--reviewer", "Dharma"])
    assert real.exit_code == 0 and "1 applied" in real.output
    again = runner.invoke(app, ["review", "import", str(run_json), str(out / "report.xlsx"), "--slot", "1", "--reviewer", "Dharma"])
    assert "1 unchanged" in again.output
