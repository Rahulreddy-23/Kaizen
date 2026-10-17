"""Quick view inherits run access, live review status and read-only guarantees."""

import pymupdf
import pytest
from fastapi.testclient import TestClient

from kaizen.api.app import RunCache, create_app
from kaizen.ingest.hashing import sha256_file
from kaizen.review.access import AccessStore
from kaizen.review.sessions import UserStore
from kaizen.workspace import Workspace
from tests.unit.test_quick_view import fixture


@pytest.fixture(scope="module")
def env(tmp_path_factory):
    root = tmp_path_factory.mktemp("quick-view")
    ws = Workspace(root / "ws")
    access = AccessStore(ws.db)
    access.bootstrap("owner@bd.com", "test-password-long")
    client = TestClient(create_app(ws))
    client.post("/api/auth/signin", json={"email": "owner@bd.com", "password": "test-password-long"})
    run = fixture()
    for doc in run.documents:
        path = root / f"{doc.id}.pdf"
        with pymupdf.open() as pdf:
            p = pdf.new_page()
            p.insert_text((20, 40), doc.items[0].description)
            pdf.save(path)
        doc.path = str(path)
        doc.sha256 = sha256_file(path)
        for item in doc.items:
            item.evidence.file = str(path)
            item.evidence.file_sha256 = doc.sha256
    RunCache(ws).put(run)
    ws.db.conn.execute("UPDATE runs SET owner=? WHERE run_id=?", ("owner@bd.com", run.metadata.run_id))
    ws.db.conn.commit()
    yield ws, client, run
    ws.db.close()


def test_endpoint_pages_evidence_and_filters(env):
    ws, client, run = env
    base = f"/api/runs/{run.metadata.run_id}/quick-view"
    response = client.get(base)
    assert response.status_code == 200, response.text
    data = response.json()
    assert data["total"] == 1 and data["counts"]["matched"] == 1
    assert len(data["documents"]) == 3 and all(d["pages"] == 1 for d in data["documents"])
    assert client.get(base, params={"sku": "not-in-run"}).status_code == 404
    assert client.get(base, params={"limit": 0}).status_code == 422
    assert client.get(base, params={"offset": -1}).status_code == 422
    assert client.get(base, params={"view": "unsupported"}).status_code == 422
    assert client.get(base, params={"search": "unrelated"}).json()["total"] == 0
    assert client.get(base, params={"view": "text"}).json()["rows"][0]["cells"]["BOM"]["items"][0]["raw_text"]
    assert "no-store" in response.headers["cache-control"]


def test_lightweight_change_check_is_stable_for_reads_and_changes_after_edits(env):
    ws, client, run = env
    assert TestClient(create_app(ws)).get('/api/changes').status_code == 401
    first = client.get('/api/changes')
    assert first.status_code == 200 and 'no-store' in first.headers['cache-control']
    token = first.json()['revision']
    assert len(token) == 64 and client.get('/api/changes').json()['revision'] == token
    base = f'/api/runs/{run.metadata.run_id}'
    assert client.patch(base, json={'name': 'Renamed for change test'}).status_code == 200
    edited = client.get('/api/changes').json()['revision']
    assert edited != token and client.get('/api/changes').json()['revision'] == edited
    with ws.db.transaction():
        ws.db.conn.execute("INSERT INTO learning_jobs (run_id,generation,state,requested_by,updated_at) VALUES (?,1,'queued','test','now')", (run.metadata.run_id,))
    queued = client.get('/api/changes').json()['revision']
    assert queued != edited
    with ws.db.transaction():
        ws.db.conn.execute("UPDATE learning_jobs SET state='complete',completed_generation=1 WHERE run_id=?", (run.metadata.run_id,))
    assert client.get('/api/changes').json()['revision'] != queued


def test_private_page_cache_reuses_render_but_rechecks_access_and_source(env, monkeypatch):
    ws, client, run = env
    doc = run.documents[0]
    url = f'/api/runs/{run.metadata.run_id}/documents/{doc.id}/pages/1'
    original = pymupdf.Page.get_pixmap
    renders = []

    def measured(page, *args, **kwargs):
        renders.append(1)
        return original(page, *args, **kwargs)

    monkeypatch.setattr(pymupdf.Page, 'get_pixmap', measured)
    one = client.get(url, params={'dpi':150, 'format':'jpeg'})
    two = client.get(url, params={'dpi':150, 'format':'jpeg'})
    assert one.status_code == two.status_code == 200
    assert one.content == two.content and len(renders) == 1
    assert one.headers['content-type'] == 'image/jpeg' and 'no-store' in one.headers['cache-control']
    assert client.get(url, params={'dpi':150}).headers['content-type'] == 'image/png'
    assert len(renders) == 2
    assert TestClient(create_app(ws)).get(url).status_code == 401
    assert client.get(url, params={'format':'unsafe'}).status_code == 422
    from pathlib import Path
    # A changed source must never receive a previous bitmap under the same URL.
    path = Path(doc.path)
    before = path.stat()
    import os
    os.utime(path, ns=(before.st_atime_ns,before.st_mtime_ns + 1_000_000))
    assert client.get(url, params={'dpi':150,'format':'jpeg'}).status_code == 200
    assert len(renders) == 3
    other = TestClient(create_app(ws))
    UserStore(ws.db).create('cached-viewer@bd.com','other-test-password')
    access = AccessStore(ws.db)
    access.register('cached-viewer@bd.com')
    access.update('cached-viewer@bd.com','owner@bd.com','approved',False)
    other.post('/api/auth/signin',json={'email':'cached-viewer@bd.com','password':'other-test-password'})
    access.share(run.metadata.run_id,'cached-viewer@bd.com','view','owner@bd.com')
    assert other.get(url, params={'dpi':150,'format':'jpeg'}).status_code == 200
    access.share(run.metadata.run_id,'cached-viewer@bd.com',None,'owner@bd.com')
    assert other.get(url, params={'dpi':150,'format':'jpeg'}).status_code == 404


def test_only_signed_in_users_with_run_access_can_read_quick_view(env):
    ws, owner, run = env
    base = f"/api/runs/{run.metadata.run_id}/quick-view"
    other = TestClient(create_app(ws))
    assert other.get(base).status_code == 401
    UserStore(ws.db).create("viewer@bd.com", "other-test-password")
    AccessStore(ws.db).register("viewer@bd.com")
    AccessStore(ws.db).update("viewer@bd.com", "owner@bd.com", "approved", False)
    assert other.post("/api/auth/signin", json={"email": "viewer@bd.com", "password": "other-test-password"}).status_code == 200
    assert other.get(base).status_code == 404
    AccessStore(ws.db).share(run.metadata.run_id, "viewer@bd.com", "view", "owner@bd.com")
    assert other.get(base).status_code == 200
    before = ws.db.conn.execute("SELECT COUNT(*) FROM decisions").fetchone()[0]
    assert other.get(base, params={"view": "text"}).status_code == 200
    assert ws.db.conn.execute("SELECT COUNT(*) FROM decisions").fetchone()[0] == before
    assert other.post(f"/api/runs/{run.metadata.run_id}/bulk-accept", json={}).status_code == 403


def test_decision_changes_are_visible_without_reparsing(env):
    ws, client, run = env
    base = f"/api/runs/{run.metadata.run_id}"
    detail = client.get(f"{base}/results/pair0").json()
    saved = client.post(f"{base}/decisions", json={"row_id": "pair0", "decision": "NEEDS_MORE_INFORMATION", "expected_revision": detail["revision"]})
    assert saved.status_code == 200, saved.text
    row = client.get(f"{base}/quick-view").json()["rows"][0]
    assert row["status"] == "attention" and row["cells"]["LABEL"]["status"] == "attention"
    assert any(c["state"] == "REVIEWED" for c in row["checks"])


def test_quick_export_uses_current_snapshot_and_private_download(env):
    from io import BytesIO

    from openpyxl import load_workbook

    ws, client, run = env
    base = f"/api/runs/{run.metadata.run_id}/quick-view"
    before = ws.db.conn.execute("SELECT COUNT(*) FROM decisions").fetchone()[0]
    projection = client.get(base).json()
    assert projection["can_export"] is True
    response = client.get(f"{base}/export.xlsx", params={"sku": "SKU-A"})
    assert response.status_code == 200, response.text
    assert "no-store" in response.headers["cache-control"]
    assert "attachment" in response.headers["content-disposition"]
    book = load_workbook(BytesIO(response.content))
    expected = "Matched" if projection["rows"][0]["status"] == "matched" else "Missing or needs review"
    assert book["1 SKU-A"]["A2"].value == expected
    assert book["Sources"].max_row == 4
    assert client.get(f"{base}/export.xlsx").status_code == 200
    assert client.get(f"{base}/export.xlsx", params={"sku": "not-in-run"}).status_code == 404
    assert ws.db.conn.execute("SELECT COUNT(*) FROM decisions").fetchone()[0] == before
    other = TestClient(create_app(ws))
    assert other.get(f"{base}/export.xlsx").status_code == 401
    assert other.post("/api/auth/signin", json={"email": "viewer@bd.com", "password": "other-test-password"}).status_code == 200
    assert other.get(base).json()["can_export"] is False
    assert other.get(f"{base}/export.xlsx").status_code == 403
