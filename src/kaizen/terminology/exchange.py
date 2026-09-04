"""Excel / CSV exchange for relationships so quality teams can maintain terminology in the tool they already use."""

import csv
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook, load_workbook
from openpyxl.styles import Font, PatternFill
from pydantic import ValidationError

from kaizen.models import DocType, Relationship
from kaizen.terminology.repository import TerminologyRepository

COLUMNS = ["ID", "Canonical", "Aliases", "Scope", "Doc Types", "Item Anchors", "Provenance", "Created By", "Active", "Notes", "Version", "Created At", "Updated At"]
_SEP = "; "


@dataclass
class ImportResult:
    created: int = 0
    updated: int = 0
    unchanged: int = 0
    errors: list[str] = field(default_factory=list)

    def summary(self) -> str:
        return f"created {self.created}, updated {self.updated}, unchanged {self.unchanged}, errors {len(self.errors)}"


def _rel_to_row(r: Relationship) -> list[Any]:
    return [r.id, r.canonical, _SEP.join(r.aliases), r.scope, _SEP.join(d.value for d in r.doc_types), _SEP.join(r.item_anchors), r.provenance, r.created_by, "Y" if r.active else "N", r.notes, r.version, r.created_at.isoformat(timespec="seconds"), r.updated_at.isoformat(timespec="seconds")]


def _split(v: Any) -> list[str]:
    if v is None:
        return []
    return [x.strip() for x in str(v).replace(";", "\n").replace(",", "\n").split("\n") if x.strip()] if "\n" in str(v) or ";" in str(v) else [x.strip() for x in str(v).split(",") if x.strip()]


def _split_aliases(v: Any) -> list[str]:
    # aliases may legitimately contain commas ("TAPE, SURGICAL"); only ';' or newlines separate entries
    if v is None:
        return []
    return [x.strip() for x in str(v).replace("\n", ";").split(";") if x.strip()]


def _truthy(v: Any) -> bool:
    return str(v).strip().upper() in ("Y", "YES", "TRUE", "1", "ACTIVE")


def export_xlsx(repo: TerminologyRepository, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    ws = wb.active
    ws.title = "Relationships"
    ws.append(COLUMNS)
    for c in ws[1]:
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = PatternFill("solid", fgColor="1F3864")
    for r in repo.list(active_only=False):
        ws.append(_rel_to_row(r))
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:M{max(ws.max_row, 2)}"
    for col, width in zip("ABCDEFGHIJKLM", (10, 40, 60, 18, 16, 16, 12, 14, 8, 50, 9, 22, 22)):
        ws.column_dimensions[col].width = width
    info = wb.create_sheet("How to edit")
    for line in (
        "Edit rows on the Relationships sheet and import the file with `kaizen terminology import <file>`.",
        "Leave ID blank to create a new relationship; keep the ID to update an existing one (a new version is recorded).",
        "Aliases and Item Anchors are separated by semicolons. Scope is 'global', 'family:<prefix>' or 'sku:<code>'.",
        "Doc Types (optional): BOM, LABEL, DRAWING, PCO separated by semicolons; blank means all document types.",
        "Active: Y or N. Version, Created At and Updated At are informational and ignored on import.",
    ):
        info.append([line])
    info.column_dimensions["A"].width = 120
    wb.save(path)
    return path


def export_csv(repo: TerminologyRepository, path: Path | str) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    with open(path, "w", newline="", encoding="utf-8") as fh:
        w = csv.writer(fh)
        w.writerow(COLUMNS)
        for r in repo.list(active_only=False):
            w.writerow(_rel_to_row(r))
    return path


def _apply_rows(repo: TerminologyRepository, rows: list[dict[str, Any]], imported_by: str, source: str) -> ImportResult:
    result = ImportResult()
    for n, row in enumerate(rows, start=2):
        canonical = (row.get("Canonical") or "").strip() if row.get("Canonical") is not None else ""
        if not canonical:
            continue
        rel_id = (str(row.get("ID")).strip() if row.get("ID") not in (None, "") else "")
        fields = {
            "canonical": canonical,
            "aliases": _split_aliases(row.get("Aliases")),
            "scope": (str(row.get("Scope") or "global")).strip() or "global",
            "doc_types": [d.upper() for d in _split(row.get("Doc Types"))],
            "item_anchors": _split(row.get("Item Anchors")),
            "notes": str(row.get("Notes") or ""),
            "active": _truthy(row.get("Active")) if row.get("Active") not in (None, "") else True,
        }
        try:
            [DocType(d) for d in fields["doc_types"]]
            Relationship(id="REL-000", **fields)  # validates scope and shapes
        except (ValidationError, ValueError) as e:
            msg = str(e).splitlines()
            result.errors.append(f"row {n}: " + next((m.strip() for m in msg if "scope" in m.lower() or "doc" in m.lower() or "Value error" in m), msg[0]))
            continue
        provenance = str(row.get("Provenance") or "imported").strip() or "imported"
        current = repo.get(rel_id) if rel_id else None
        if current is None:
            repo.create(rel_id=rel_id or None, provenance=provenance, created_by=imported_by, **fields)
            result.created += 1
            continue
        changed = {k: v for k, v in fields.items() if getattr(current, k) != (v if k != "doc_types" else [DocType(d) for d in v])}
        if not changed:
            result.unchanged += 1
            continue
        repo.update(current.id, changed_by=imported_by, change_note=f"imported from {source}", **changed)
        result.updated += 1
    return result


def import_xlsx(repo: TerminologyRepository, path: Path | str, imported_by: str = "import") -> ImportResult:
    wb = load_workbook(path, data_only=True)
    ws = wb["Relationships"] if "Relationships" in wb.sheetnames else wb.worksheets[0]
    header = [c.value for c in ws[1]]
    rows = [dict(zip(header, r)) for r in ws.iter_rows(min_row=2, values_only=True)]
    return _apply_rows(repo, rows, imported_by, Path(path).name)


def import_csv(repo: TerminologyRepository, path: Path | str, imported_by: str = "import") -> ImportResult:
    with open(path, newline="", encoding="utf-8-sig") as fh:
        rows = list(csv.DictReader(fh))
    return _apply_rows(repo, rows, imported_by, Path(path).name)
