"""Parser for PCO change forms (FM00835 'Bill of Material and Routing Form') in XLSX/CSV or PDF form.

A PCO describes INTENDED changes: affected codes (SKUs) plus rows of item-number actual/proposed, description,
quantity and operation-sequence actual/proposed. Each row is classified as ADD, DELETE, SUBSTITUTE or MODIFY.
"""

import csv
import re
from decimal import Decimal
from pathlib import Path
from typing import Any

import openpyxl
import pymupdf

from kaizen.ingest.hashing import document_id, sha256_file
from kaizen.ingest.pdf_words import ColumnBand, assign_columns, extract_words, group_lines, line_bbox, line_text
from kaizen.ingest.quantity import parse_decimal
from kaizen.models import BBox, DocType, Document, DocumentItem, Evidence, ItemCategory

PARSER_NAME = "pco"
PARSER_VERSION = "1"

COLUMN_ALIASES = {
    "affected_codes": {"affected codes", "affected code", "affected skus", "affected items"},
    "item_actual": {"item number actual", "item actual", "actual item number", "item no actual"},
    "item_proposed": {"item number proposed", "item proposed", "proposed item number", "item no proposed"},
    "description": {"description", "item description"},
    "qty_actual": {"quantity actual", "qty actual"},
    "qty_proposed": {"quantity proposed", "qty proposed"},
    "seq_actual": {"oper seq# actual", "oper seq actual", "operation sequence actual", "oper seq # actual"},
    "seq_proposed": {"oper seq# proposed", "oper seq proposed", "operation sequence proposed", "oper seq # proposed"},
    "scrap_actual": {"percent scrap actual", "scrap actual"},
    "scrap_proposed": {"percent scrap proposed", "scrap proposed"},
}
_CODE_RE = re.compile(r"^[A-Z0-9]{5,14}$")
_PCO_RE = re.compile(r"(PCO\s*\d+)\s*,\s*(\d+)", re.IGNORECASE)
_BRANCH_RE = re.compile(r"Branch:?\s*([A-Z0-9]+)", re.IGNORECASE)
_FORM_RE = re.compile(r"\b(FM\d{5})\b")
_FORM_REV_RE = re.compile(r"Revision\s+(\d+)", re.IGNORECASE)
_DELETE_WORDS = {"DELETE", "DEL", "REMOVE", "REMOVED", "DELETED"}


def change_kind(actual: str, proposed: str) -> tuple[str, str]:
    a, p = actual.strip(), proposed.strip()
    if p.upper() in _DELETE_WORDS and a:
        return "DELETE", f"DELETE:{a}"
    if not a and p:
        return "ADD", f"ADD:{p}"
    if a and p and a != p:
        return "SUBSTITUTE", f"SUBSTITUTE:{a}>{p}"
    if a and p and a == p:
        return "MODIFY", f"MODIFY:{a}"
    if a and not p:
        return "MODIFY", f"MODIFY:{a}"
    return "UNKNOWN", "UNKNOWN"


def _header_fields(text: str) -> dict[str, Any]:
    header: dict[str, Any] = {}
    if m := _PCO_RE.search(text):
        header["pco_number"] = m.group(1).replace(" ", "").upper()
        header["revision"] = m.group(2)
    if m := _BRANCH_RE.search(text):
        header["branch"] = m.group(1)
    if m := _FORM_RE.search(text):
        header["form"] = m.group(1)
    if m := _FORM_REV_RE.search(text):
        header["form_revision"] = m.group(1)
    for kind in ("New", "Change", "Substitute"):
        if re.search(r"\[\s*X\s*\]\s*" + kind, text, re.IGNORECASE):
            header["change_type"] = kind
    for bill in ("ALT", "M"):
        if re.search(r"\[\s*X\s*\]\s*" + bill + r"\b", text):
            header["bill_type"] = bill
            break
    return header


def _decimal(s: str) -> Decimal | None:
    return parse_decimal(s)


class _Builder:
    def __init__(self, doc_id: str, sha: str, path: Path):
        self.doc_id, self.sha, self.path = doc_id, sha, path
        self.items: list[DocumentItem] = []
        self.codes: list[str] = []
        self.warnings: list[str] = []

    def add_code(self, code: str, raw: str, locator: str, page: int | None = None, bbox: BBox | None = None) -> None:
        code = code.strip()
        if not _CODE_RE.match(code):
            self.warnings.append(f"{locator}: '{code}' does not look like an affected code; ignored")
            return
        self.codes.append(code)
        self.items.append(
            DocumentItem(
                id=f"{self.doc_id}:code{len(self.codes)}", doc_id=self.doc_id, doc_type=DocType.PCO, item_number=code, description=f"affected code {code}",
                category=ItemCategory.ADMINISTRATIVE, category_reason="PCO affected code", attributes={"kind": "affected_code"},
                evidence=Evidence(file=str(self.path), file_sha256=self.sha, page=page, bbox=bbox, raw_text=raw, locator=locator, extraction_method="pdf_text" if page else "table"),
            )
        )

    def add_change(self, cells: dict[str, str], raw: str, locator: str, page: int | None = None, bbox: BBox | None = None) -> None:
        actual, proposed = cells.get("item_actual", "").strip(), cells.get("item_proposed", "").strip()
        desc = cells.get("description", "").strip()
        if not (actual or proposed or desc):
            return
        kind, key = change_kind(actual, proposed)
        if kind == "UNKNOWN":
            self.warnings.append(f"{locator}: change row without item numbers ignored ('{desc}')")
            return
        n = sum(1 for i in self.items if i.attributes.get("kind") == "change") + 1
        item_number = actual if kind == "DELETE" else (proposed or actual)
        qty = _decimal(cells.get("qty_proposed", ""))
        self.items.append(
            DocumentItem(
                id=f"{self.doc_id}:chg{n}", doc_id=self.doc_id, doc_type=DocType.PCO, item_number=item_number, description=desc, quantity=qty,
                oper_seq=cells.get("seq_proposed", "").strip() or None, category=ItemCategory.ADMINISTRATIVE, category_reason="PCO change instruction",
                attributes={
                    "kind": "change", "change_kind": kind, "change_key": key, "item_actual": actual, "item_proposed": "" if kind == "DELETE" else proposed,
                    "qty_actual": cells.get("qty_actual", "").strip(), "qty_proposed": cells.get("qty_proposed", "").strip(),
                    "seq_actual": cells.get("seq_actual", "").strip(), "seq_proposed": cells.get("seq_proposed", "").strip(),
                    "scrap_actual": cells.get("scrap_actual", "").strip(), "scrap_proposed": cells.get("scrap_proposed", "").strip(),
                },
                evidence=Evidence(file=str(self.path), file_sha256=self.sha, page=page, bbox=bbox, raw_text=raw, locator=locator, extraction_method="pdf_text" if page else "table"),
            )
        )


def _column_map(row: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, text in enumerate(row):
        key = text.strip().lower().replace("  ", " ")
        for name, aliases in COLUMN_ALIASES.items():
            if key in aliases and name not in mapping:
                mapping[name] = idx
    return mapping if "item_proposed" in mapping or "affected_codes" in mapping else None


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, float) and v.is_integer():
        return str(int(v))
    return str(v).strip()


def _parse_table(path: Path, b: _Builder, header: dict[str, Any]) -> None:
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            sheets = [(path.name, [[c.strip() for c in row] for row in csv.reader(fh)])]
    else:
        wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
        sheets = [(ws.title, [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]) for ws in wb.worksheets]
        wb.close()
    for sheet, rows in sheets:
        header.update(_header_fields("\n".join(" ".join(r) for r in rows[:40])))
        header_idx, colmap = None, None
        for idx, row in enumerate(rows):
            colmap = _column_map(row)
            if colmap:
                header_idx = idx
                break
        if header_idx is None:
            continue
        def get(row, name, colmap=colmap):
            return row[colmap[name]] if name in colmap and colmap[name] < len(row) else ""
        for r_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            locator = f"sheet '{sheet}', row {r_idx}"
            code = get(row, "affected_codes")
            if code:
                b.add_code(code, code, locator)
            cells = {name: get(row, name) for name in COLUMN_ALIASES if name != "affected_codes"}
            b.add_change(cells, " | ".join(c for c in row if c), locator)
        return


_PDF_LABELS = [("affected_codes", ("Affected",)), ("item_actual", ("Item",)), ("item_proposed", ("Item",)), ("description", ("Description",)), ("qty_actual", ("Quantity",)), ("qty_proposed", ("Quantity",)), ("seq_actual", ("Oper",)), ("seq_proposed", ("Oper",)), ("scrap_actual", ("Percent",)), ("scrap_proposed", ("Percent",))]


def _parse_pdf(path: Path, b: _Builder, header: dict[str, Any]) -> None:
    pdf = pymupdf.open(path)
    for page in pdf:
        page_no = page.number + 1
        header.update({k: v for k, v in _header_fields(page.get_text()).items() if k not in header})
        lines = group_lines(extract_words(page), y_tol=2.5)
        head_idx = next((i for i, l in enumerate(lines) if "Affected" in line_text(l) and "Description" in line_text(l)), None)
        if head_idx is None:
            continue
        words = sorted(lines[head_idx], key=lambda w: w.x0)
        starts: list[tuple[str, float]] = []
        pos = 0
        for name, tokens in _PDF_LABELS:
            found = next((i for i in range(pos, len(words)) if words[i].text == tokens[0]), None)
            if found is None:
                continue
            starts.append((name, words[found].x0))
            pos = found + 1
        bands = [ColumnBand(name, x0 - 2, (starts[i + 1][1] - 2 if i + 1 < len(starts) else 10_000)) for i, (name, x0) in enumerate(starts)]
        for line in lines[head_idx + 1 :]:
            text = line_text(line)
            if text.replace("Actual", "").replace("Proposed", "").strip() == "":
                continue
            cols = assign_columns(line, bands)
            code_words = cols.get("affected_codes", [])
            if code_words:
                cw = code_words[0]
                b.add_code(cw.text, cw.text, f"page {page_no}, y {cw.y0:.0f}", page_no, BBox(x0=cw.x0, y0=cw.y0, x1=cw.x1, y1=cw.y1))
            cells = {name: " ".join(w.text for w in cols.get(name, [])) for name, _ in _PDF_LABELS if name != "affected_codes"}
            if any(cells.values()):
                change_words = [w for name in cells for w in cols.get(name, [])]
                b.add_change(cells, " ".join(w.text for w in change_words), f"page {page_no}, y {line[0].y0:.0f}", page_no, line_bbox(change_words))
    pdf.close()


def parse_pco(path: Path | str) -> Document:
    path = Path(path)
    sha = sha256_file(path)
    doc_id = document_id("pco", sha, path)
    header: dict[str, Any] = {}
    b = _Builder(doc_id, sha, path)
    if path.suffix.lower() == ".pdf":
        _parse_pdf(path, b, header)
    else:
        _parse_table(path, b, header)
    header["affected_codes"] = list(b.codes)
    if not b.codes:
        b.warnings.append("no affected codes found on the PCO")
    if not any(i.attributes.get("kind") == "change" for i in b.items):
        b.warnings.append("no change rows found on the PCO")
    if "pco_number" not in header:
        b.warnings.append("PCO number not found")
    return Document(id=doc_id, doc_type=DocType.PCO, path=str(path), sha256=sha, sku=None, header=header, items=b.items, parser_name=PARSER_NAME, parser_version=PARSER_VERSION, warnings=b.warnings)
