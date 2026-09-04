"""Parser for BOM exports in XLSX/CSV form (e.g. a JDE grid export). Header row is found by column names."""

import csv
from datetime import date, datetime
from pathlib import Path
from typing import Any

import openpyxl

from kaizen.ingest.bom_categorize import CATEGORIZER_VERSION, categorize
from kaizen.ingest.hashing import document_id, sha256_file
from kaizen.ingest.quantity import parse_decimal
from kaizen.models import DocType, Document, DocumentItem, Evidence

PARSER_NAME = "bom_table"
PARSER_VERSION = f"1+cat{CATEGORIZER_VERSION}"

COLUMN_ALIASES: dict[str, set[str]] = {
    "level": {"level"},
    "component_item": {"component item", "item number", "item", "component", "part number", "item no", "item no."},
    "component_description": {"component description", "description", "item description", "desc"},
    "branch_plant": {"branch/plant", "branch", "plant"},
    "quantity_per": {"quantity per", "qty per", "quantity", "qty"},
    "ext_qty": {"ext qty", "extended quantity", "ext quantity"},
    "um": {"um", "uom", "unit"},
    "t": {"t"},
    "effective_from": {"effective from", "from", "eff from"},
    "effective_thru": {"effective thru", "thru", "eff thru", "effective to", "to"},
    "oper_seq": {"oper seq no", "oper seq", "operation sequence", "op seq", "oper seq #", "oper seq no."},
    "flags": {"flags"},
}
HEADER_KEYS = {
    "parent item": "parent_item",
    "parent description": "parent_description",
    "branch/plant": "branch_plant",
    "as of date": "as_of_date",
    "type": "type",
    "batch quantity": "batch_quantity",
    "batch uom": "batch_uom",
    "bill revision level": "bill_revision_level",
    "requested quantity": "requested_quantity",
    "requested uom": "requested_uom",
}


def _cell(v: Any) -> str:
    if v is None:
        return ""
    if isinstance(v, datetime):
        return v.strftime("%m/%d/%y")
    if isinstance(v, date):
        return v.strftime("%m/%d/%y")
    if isinstance(v, float) and v.is_integer() and abs(v) < 1e12:
        return str(v)
    return str(v).strip()


def _column_map(row: list[str]) -> dict[str, int] | None:
    mapping: dict[str, int] = {}
    for idx, text in enumerate(row):
        key = text.strip().lower()
        if not key:
            continue
        for name, aliases in COLUMN_ALIASES.items():
            if key in aliases and name not in mapping:
                mapping[name] = idx
                break
    if "component_item" in mapping and ("component_description" in mapping or "quantity_per" in mapping):
        return mapping
    return None


def _load_sheets(path: Path) -> list[tuple[str, list[list[str]]]]:
    if path.suffix.lower() == ".csv":
        with open(path, newline="", encoding="utf-8-sig") as fh:
            return [(path.name, [[c.strip() for c in row] for row in csv.reader(fh)])]
    wb = openpyxl.load_workbook(path, data_only=True, read_only=True)
    sheets = []
    for ws in wb.worksheets:
        sheets.append((ws.title, [[_cell(c) for c in row] for row in ws.iter_rows(values_only=True)]))
    wb.close()
    return sheets


def _parse_date(s: str) -> date | None:
    for fmt in ("%m/%d/%y", "%m/%d/%Y", "%Y-%m-%d"):
        try:
            return datetime.strptime(s, fmt).date()
        except ValueError:
            continue
    return None


def parse_bom_table(path: Path | str) -> Document:
    path = Path(path)
    sha = sha256_file(path)
    doc_id = document_id("bom", sha, path)
    header: dict[str, str] = {}
    items: list[DocumentItem] = []
    warnings: list[str] = []
    for sheet_name, rows in _load_sheets(path):
        header_idx, colmap = None, None
        for idx, row in enumerate(rows):
            colmap = _column_map(row)
            if colmap:
                header_idx = idx
                break
        if header_idx is None:
            continue
        for row in rows[:header_idx]:
            for c, text in enumerate(row):
                key = HEADER_KEYS.get(text.strip().lower())
                if key and key not in header:
                    value = next((v for v in row[c + 1 :] if v), "")
                    header[key] = value
        as_of = _parse_date(header.get("as_of_date", ""))
        def get(row, name, colmap=colmap):
            return row[colmap[name]] if name in colmap and colmap[name] < len(row) else ""
        for r_idx, row in enumerate(rows[header_idx + 1 :], start=header_idx + 2):
            item_number = get(row, "component_item")
            if not item_number:
                continue
            description = get(row, "component_description")
            qty_text = get(row, "quantity_per")
            confidence = 1.0
            quantity = parse_decimal(qty_text)
            if quantity is None:
                confidence = 0.5
                warnings.append(f"sheet '{sheet_name}' row {r_idx}: quantity '{qty_text}' is not numeric")
            eff_from, eff_thru = get(row, "effective_from"), get(row, "effective_thru")
            d_from, d_thru = _parse_date(eff_from), _parse_date(eff_thru)
            is_active = not ((as_of and d_thru and d_thru < as_of) or (as_of and d_from and d_from > as_of))
            category, reason = categorize(item_number, description, quantity)
            items.append(
                DocumentItem(
                    id=f"{doc_id}:r{len(items) + 1}",
                    doc_id=doc_id,
                    doc_type=DocType.BOM,
                    sku=header.get("parent_item"),
                    item_number=item_number,
                    description=description,
                    quantity=quantity,
                    uom=get(row, "um") or None,
                    oper_seq=get(row, "oper_seq") or None,
                    attributes={
                        "level": get(row, "level"),
                        "branch_plant": get(row, "branch_plant"),
                        "ext_qty": get(row, "ext_qty"),
                        "t": get(row, "t"),
                        "effective_from": eff_from,
                        "effective_thru": eff_thru,
                        "flags": get(row, "flags"),
                        "redlines": [],
                    },
                    category=category,
                    category_reason=reason,
                    is_active=is_active,
                    extraction_confidence=confidence,
                    evidence=Evidence(
                        file=str(path),
                        file_sha256=sha,
                        page=None,
                        bbox=None,
                        raw_text=" | ".join(c for c in row if c),
                        line_index=r_idx - 1,
                        sheet=sheet_name,
                        locator=f"sheet '{sheet_name}', row {r_idx}",
                        extraction_method="table",
                    ),
                )
            )
        break  # first sheet with a header row wins
    if not items and not header:
        warnings.append("no BOM header row (Component Item / Description / Quantity) found in any sheet")
    if not header.get("parent_item"):
        warnings.append("parent item not found in header block")
    return Document(
        id=doc_id,
        doc_type=DocType.BOM,
        path=str(path),
        sha256=sha,
        sku=header.get("parent_item"),
        header=header,
        items=items,
        parser_name=PARSER_NAME,
        parser_version=PARSER_VERSION,
        warnings=warnings,
    )
