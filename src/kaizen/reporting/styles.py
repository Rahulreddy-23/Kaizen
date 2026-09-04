from openpyxl.styles import Alignment, Border, Font, PatternFill, Side

HEADER_FILL = PatternFill("solid", fgColor="1F3864")
HEADER_FONT = Font(bold=True, color="FFFFFF")
TITLE_FONT = Font(bold=True, size=14)
SECTION_FONT = Font(bold=True, size=11)
WRAP = Alignment(wrap_text=True, vertical="top")
THIN = Side(style="thin", color="BFBFBF")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

CLASS_FILLS = {
    "EXACT": PatternFill("solid", fgColor="C6EFCE"),
    "EQUIVALENT": PatternFill("solid", fgColor="E2EFDA"),
    "POTENTIAL": PatternFill("solid", fgColor="FFEB9C"),
    "MISMATCH": PatternFill("solid", fgColor="FFC7CE"),
    "MISSING": PatternFill("solid", fgColor="F8CBAD"),
}
SEVERITY_FILLS = {
    "BLOCKER": PatternFill("solid", fgColor="C00000"),
    "MAJOR": PatternFill("solid", fgColor="FFC7CE"),
    "MINOR": PatternFill("solid", fgColor="FFEB9C"),
    "INFO": PatternFill("solid", fgColor="DDEBF7"),
}
