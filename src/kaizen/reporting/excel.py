"""Audit-grade Excel workbook. If a judge asks 'where did this come from?', the workbook answers."""

from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from openpyxl import Workbook
from openpyxl.formatting.rule import CellIsRule
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from kaizen.checks.bom_label import is_comparable_bom_item
from kaizen.models import Classification, DocType, DocumentItem, Run
from kaizen.reporting.styles import BORDER, CLASS_FILLS, HEADER_FILL, HEADER_FONT, SECTION_FONT, SEVERITY_FILLS, TITLE_FONT, WRAP


@dataclass
class ReviewBundle:
    """Reviewer state to merge into the workbook. Engine recommendations are never modified by it."""

    decisions: dict[str, dict[int, Any]] = field(default_factory=dict)  # row_id → slot → Decision
    finals: dict[str, Any] = field(default_factory=dict)  # row_id → Final
    states: dict[str, str] = field(default_factory=dict)  # row_id → state
    action_items: list[Any] = field(default_factory=list)
    business: Any | None = None


def _review_cells(review: ReviewBundle | None, r) -> list[Any]:
    if review is None:
        return [r.reviewer_decision, r.reviewer_comment, None, None, None, None, None, None, "ENGINE_RECOMMENDED", "", r.final_status]
    d = review.decisions.get(r.row_id, {})
    d1, d2 = d.get(1), d.get(2)
    final = review.finals.get(r.row_id)
    items = "; ".join(a.id for a in review.action_items if a.row_id == r.row_id)
    return [
        d1.decision if d1 else None, d1.comment if d1 else None, d1.reviewer if d1 else None, d1.decided_at if d1 else None,
        d2.decision if d2 else None, d2.comment if d2 else None, d2.reviewer if d2 else None, d2.decided_at if d2 else None,
        review.states.get(r.row_id, "ENGINE_RECOMMENDED"), items, f"FINALIZED: {final.final_decision}" if final else "OPEN",
    ]


def export_with_review(workspace, run: Run, path: Path | str, metrics: dict[str, Any] | None = None) -> Path:
    """Workbook with the workspace's reviewer decisions, action items and business case merged in."""
    from kaizen.review.action_items import ActionItemStore
    from kaizen.review.business import business_case
    from kaizen.review.store import ReviewStore

    review = ReviewStore(workspace.db)
    rid = run.metadata.run_id
    decisions = review.all_decisions(rid)
    finals = review.finals(rid)
    states = {r.row_id: ReviewStore.state_of(decisions.get(r.row_id, {}), finals.get(r.row_id)) for r in run.results}
    bundle = ReviewBundle(decisions=decisions, finals=finals, states=states, action_items=ActionItemStore(workspace.db).for_run(rid), business=business_case(run))
    return write_report(run, path, metrics=metrics, review=bundle)

BOM_LABEL_COLUMNS = [
    "Row ID", "SKU", "Check",
    "Source A File", "Source A Page", "Source A Locator", "Source A Item", "Source A Description", "Source A Quantity", "Source A Category",
    "Source B File", "Source B Page", "Source B Locator", "Source B Item", "Source B Description", "Source B Quantity",
    "Normalized A", "Normalized B", "Classification", "Match Level", "Score", "Relationship ID",
    "Discrepancy Type", "Discrepancy Detail", "Severity", "Recommended Action", "Requires Validation", "Explanation",
    "Reviewer Decision", "Reviewer Comment", "Reviewer Name", "Reviewer Date",
    "Reviewer 2 Decision", "Reviewer 2 Comment", "Reviewer 2 Name", "Reviewer 2 Date",
    "Review State", "Action Item", "Final Status",
]
PCO_BOM_COLUMNS = [
    "Row ID", "SKU", "Check", "PCO", "Change", "Kind", "Item Actual", "Item Proposed", "PCO Description", "Qty Proposed", "Seq Proposed", "PCO Locator",
    "BOM File", "BOM Page", "BOM Locator", "BOM Item", "BOM Description", "BOM Quantity", "BOM Oper Seq",
    "Classification", "Discrepancy Type", "Discrepancy Detail", "Severity", "Recommended Action", "Requires Validation", "Explanation",
    "Reviewer Decision", "Reviewer Comment", "Reviewer Name", "Reviewer Date",
    "Reviewer 2 Decision", "Reviewer 2 Comment", "Reviewer 2 Name", "Reviewer 2 Date",
    "Review State", "Action Item", "Final Status",
]
ACTION_ITEM_COLUMNS = ["ID", "Run", "Row ID", "SKU", "Check", "Discrepancy", "Severity", "Detail", "Recommended Action", "Owner", "Status", "Reviewer", "Created", "Updated", "Resolved In Run", "Resolved At", "Comparison Key"]
MANUAL_CHECKS = [
    ("Confirm dot sticker location per approved drawing", "Brief, cross-check activity 6 — visual check, not automated"),
    ("Confirm case label artwork against the approved redline", "Case labels are not parsed in this build"),
    ("Confirm the printed IFU revision matches the released IFU", "IFU documents are BOM lines only"),
]
COVERAGE_COLUMNS = ["Kind", "SKU / Affected Code", "Status", "Detail", "Source"]
DECISIONS = "ACCEPT,OVERRIDE,CONFIRM_DISCREPANCY,NEEDS_MORE_INFORMATION"
_WIDTHS = {"Explanation": 70, "Discrepancy Detail": 60, "Recommended Action": 45, "Source A Description": 32, "Source B Description": 40, "Normalized A": 28, "Normalized B": 32, "Reviewer Comment": 30, "Reviewer 2 Comment": 30}


def _display_path(file: str | None, root: str) -> str:
    if not file:
        return ""
    try:
        return Path(file).resolve().relative_to(Path(root).resolve()).as_posix()
    except ValueError:
        return Path(file).name


def _qty(item: DocumentItem | None) -> Any:
    if item is None or item.quantity is None:
        return None
    return float(item.quantity)


def _write_header(ws, row: int, headers: list[str]) -> None:
    for c, h in enumerate(headers, start=1):
        cell = ws.cell(row=row, column=c, value=h)
        cell.fill, cell.font, cell.border = HEADER_FILL, HEADER_FONT, BORDER


def _autosize(ws, headers: list[str], default: int = 16) -> None:
    for c, h in enumerate(headers, start=1):
        ws.column_dimensions[get_column_letter(c)].width = _WIDTHS.get(h, max(default, min(len(h) + 2, 40)))


def write_report(run: Run, path: Path | str, metrics: dict[str, Any] | None = None, review: ReviewBundle | None = None) -> Path:
    path = Path(path)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb = Workbook()
    _summary_sheet(wb.active, run, review)
    _pairing_sheet(wb.create_sheet("BOM_Label"), run, "BOM_LABEL", review)
    _pairing_sheet(wb.create_sheet("BOM_Drawing"), run, "BOM_DRAWING", review)
    _pairing_sheet(wb.create_sheet("Label_Drawing"), run, "LABEL_DRAWING", review)
    _pco_bom_sheet(wb.create_sheet("PCO_BOM"), run, review)
    _pairing_sheet(wb.create_sheet("Label_Revision"), run, "LABEL_REVISION", review)
    _action_items_sheet(wb.create_sheet("Action_Items"), run, review)
    _coverage_sheet(wb.create_sheet("Coverage"), run)
    _manual_checklist_sheet(wb.create_sheet("Manual_Checklist"), run)
    _documents_sheet(wb.create_sheet("Documents"), run)
    _relationships_sheet(wb.create_sheet("Relationships_Used"), run)
    _metadata_sheet(wb.create_sheet("Run_Metadata"), run)
    _audit_sheet(wb.create_sheet("Audit_Log"), run)
    if metrics:
        _accuracy_sheet(wb.create_sheet("Accuracy"), metrics)
    wb.save(path)
    return path


def _action_items_sheet(ws, run: Run, review: ReviewBundle | None) -> None:
    _write_header(ws, 1, ACTION_ITEM_COLUMNS)
    items = review.action_items if review else []
    for i, a in enumerate(items, start=2):
        for c, v in enumerate([a.id, a.run_id, a.row_id, a.sku, a.check_type, a.discrepancy_type, a.severity, a.detail, a.recommended_action, a.owner, a.status, a.reviewer, a.created_at, a.updated_at, a.resolved_in_run, a.resolved_at, a.comparison_key], start=1):
            ws.cell(row=i, column=c, value=v).border = BORDER
        ws.cell(row=i, column=7).fill = SEVERITY_FILLS.get(a.severity, SEVERITY_FILLS["INFO"])
    if not items:
        ws.cell(row=2, column=1, value="No action items yet — created by reviewers from confirmed discrepancies (CONFIRM_DISCREPANCY).")
    ws.freeze_panes = "A2"
    _autosize(ws, ACTION_ITEM_COLUMNS, default=14)
    ws.column_dimensions["H"].width = 70
    ws.column_dimensions["I"].width = 50


def _manual_checklist_sheet(ws, run: Run) -> None:
    headers = ["SKU", "Manual check", "Basis", "Reviewer", "Date", "Result (PASS / FAIL / N/A)", "Notes"]
    _write_header(ws, 1, headers)
    row = 1
    for g in run.groups:
        for check, basis in MANUAL_CHECKS:
            row += 1
            for c, v in enumerate([g.sku, check, basis, None, None, None, None], start=1):
                ws.cell(row=row, column=c, value=v).border = BORDER
    dv = DataValidation(type="list", formula1='"PASS,FAIL,N/A"', allow_blank=True)
    ws.add_data_validation(dv)
    dv.add(f"F2:F{max(row, 2)}")
    ws.freeze_panes = "A2"
    ws.column_dimensions["B"].width = 60
    ws.column_dimensions["C"].width = 60
    ws.column_dimensions["F"].width = 22


def _summary_sheet(ws, run: Run, review: ReviewBundle | None = None) -> None:
    ws.title = "Summary"
    m = run.metadata
    ws["A1"] = "Kaizen Cross-Check — BOM ↔ Label cross-check report"
    ws["A1"].font = TITLE_FONT
    info = [("Run ID", m.run_id), ("Generated", m.timestamp.isoformat(timespec="seconds")), ("Tool version", m.tool_version), ("Input root", m.input_root), ("Terminology version", m.terminology_version), ("Relationships in dictionary", m.terminology_count)]
    for i, (k, v) in enumerate(info, start=3):
        ws.cell(row=i, column=1, value=k).font = SECTION_FONT
        ws.cell(row=i, column=2, value=v)
    row = 3 + len(info) + 1
    ws.cell(row=row, column=1, value="Per SKU").font = SECTION_FONT
    row += 1
    headers = ["SKU", "Family", "BOM File", "Label File", "Rows", "EXACT", "EQUIVALENT", "POTENTIAL", "MISMATCH", "MISSING", "Needs Validation", "Auto-cleared", "Blockers", "Status", "BOM↔Label", "BOM↔Drawing", "Label↔Drawing", "PCO↔BOM", "Old↔New"]
    _write_header(ws, row, headers)
    docs = {d.id: d for d in run.documents}
    totals = [0] * 9
    groups_by_sku = {g.sku: g for g in run.groups}
    skus = list(groups_by_sku) + sorted({r.sku for r in run.results} - set(groups_by_sku))
    for sku in skus:
        g = groups_by_sku.get(sku)
        rows = [r for r in run.results if r.sku == sku and r.role != "exempt"]
        counts = {c.value: sum(1 for r in rows if r.classification is c) for c in Classification}
        needs = sum(1 for r in rows if r.requires_validation)
        blockers = sum(1 for r in rows for d in r.discrepancies if d.severity.value == "BLOCKER")
        bom = next((docs[i] for i in g.document_ids if docs[i].doc_type is DocType.BOM), None) if g else None
        label = next((docs[i] for i in g.document_ids if docs[i].doc_type is DocType.LABEL), None) if g else None
        status = "BLOCKED" if blockers else "NEEDS REVIEW" if needs else "AUTO-CLEARED" if rows else "NOT CHECKED"
        per_check = {ct: sum(1 for r in rows if r.check.value == ct) for ct in ("BOM_LABEL", "BOM_DRAWING", "LABEL_DRAWING", "PCO_BOM", "LABEL_REVISION")}
        values = [sku, g.family if g else "", _display_path(bom.path, m.input_root) if bom else "", _display_path(label.path, m.input_root) if label else "", len(rows), counts["EXACT"], counts["EQUIVALENT"], counts["POTENTIAL"], counts["MISMATCH"], counts["MISSING"], needs, len(rows) - needs, blockers, status, per_check["BOM_LABEL"], per_check["BOM_DRAWING"], per_check["LABEL_DRAWING"], per_check["PCO_BOM"], per_check["LABEL_REVISION"]]
        row += 1
        for c, v in enumerate(values, start=1):
            ws.cell(row=row, column=c, value=v).border = BORDER
        for k in range(9):
            totals[k] += values[4 + k]
    row += 1
    ws.cell(row=row, column=1, value="TOTAL").font = SECTION_FONT
    for k, v in enumerate(totals):
        ws.cell(row=row, column=5 + k, value=v).font = SECTION_FONT
    row += 2
    ws.cell(row=row, column=1, value="Effort view").font = SECTION_FONT
    total_rows = totals[0]
    cleared = totals[7]
    ws.cell(row=row + 1, column=1, value="Rows auto-cleared (exact/equivalent, no discrepancy)")
    ws.cell(row=row + 1, column=2, value=cleared)
    ws.cell(row=row + 2, column=1, value="Rows needing reviewer validation")
    ws.cell(row=row + 2, column=2, value=totals[6])
    ws.cell(row=row + 3, column=1, value="Auto-cleared share")
    ws.cell(row=row + 3, column=2, value=(cleared / total_rows) if total_rows else 0)
    ws.cell(row=row + 3, column=2).number_format = "0%"
    row += 5
    if review is not None:
        ws.cell(row=row, column=1, value="Review status (rows by state)").font = SECTION_FONT
        counts: dict[str, int] = {s: 0 for s in ("ENGINE_RECOMMENDED", "REVIEWER_1_COMPLETE", "REVIEWER_2_COMPLETE", "AGREED", "DISAGREEMENT", "FINALIZED")}
        for st in review.states.values():
            counts[st] = counts.get(st, 0) + 1
        for i, (k, v) in enumerate(counts.items(), start=1):
            ws.cell(row=row + i, column=1, value=k)
            ws.cell(row=row + i, column=2, value=v)
        row += len(counts) + 2
        if review.business is not None:
            b = review.business
            ws.cell(row=row, column=1, value="Estimated review effort (from this run; assumptions listed)").font = SECTION_FONT
            lines = [("Estimated minutes per SKU (both reviewers' rows)", b.estimated_minutes_per_sku), ("Estimated minutes saved per SKU vs 60-minute baseline", b.minutes_saved_per_sku), ("Estimated reduction (measured on this run)", f"{b.reduction_pct}%"), ("Estimated hours saved per project", b.hours_saved_per_project), ("Estimated annual savings (USD)", b.annual_savings), ("Meets 50% target (measured)", "yes" if b.meets_target else "NO"), ("Projection after confirming strong POTENTIAL pairings as relationships: rows still needing validation", b.needs_validation_after_confirmation), ("Projection: estimated reduction", f"{b.reduction_pct_after_confirmation}%"), ("Projection: annual savings (USD)", b.annual_savings_after_confirmation), ("Projection: meets 50% target", "yes" if b.meets_target_after_confirmation else "NO")]
            for i, (k, v) in enumerate(lines, start=1):
                ws.cell(row=row + i, column=1, value=k)
                ws.cell(row=row + i, column=2, value=v)
            for j, a_ in enumerate(b.assumptions, start=len(lines) + 1):
                ws.cell(row=row + j, column=1, value=f"assumption: {a_}")
            row += len(lines) + len(b.assumptions) + 2
    ws.cell(row=row, column=1, value="Coverage & parser warnings").font = SECTION_FONT
    warnings = [f"COVERAGE {c.kind} {c.sku}: {c.status} — {c.detail}" for c in run.coverage if c.status != "OK"] + list(run.warnings) + [f"{g.sku}: {w}" for g in run.groups for w in g.warnings]
    if not warnings:
        ws.cell(row=row + 1, column=1, value="none")
    for i, w in enumerate(warnings, start=1):
        ws.cell(row=row + i, column=1, value=w)
    row += len(warnings) + 2
    ws.cell(row=row, column=1, value="Capabilities in this build").font = SECTION_FONT
    for i, (k, v) in enumerate(m.capabilities.items(), start=1):
        ws.cell(row=row + i, column=1, value=k)
        ws.cell(row=row + i, column=2, value=v)
    ws.column_dimensions["A"].width = 58
    ws.column_dimensions["B"].width = 40
    for c in range(3, 15):
        ws.column_dimensions[get_column_letter(c)].width = 14


def _finish_results_sheet(ws, columns: list[str], n_rows: int) -> None:
    last = max(n_rows + 1, 2)
    cls_col = get_column_letter(columns.index("Classification") + 1)
    for value, fill in CLASS_FILLS.items():
        ws.conditional_formatting.add(f"{cls_col}2:{cls_col}{last}", CellIsRule(operator="equal", formula=[f'"{value}"'], fill=fill))
    dv = DataValidation(type="list", formula1=f'"{DECISIONS}"', allow_blank=True)
    ws.add_data_validation(dv)
    for name in ("Reviewer Decision", "Reviewer 2 Decision"):
        col = get_column_letter(columns.index(name) + 1)
        dv.add(f"{col}2:{col}{last}")
    status = DataValidation(type="list", formula1='"OPEN,ACCEPTED,ACTION_REQUIRED,CLOSED"', allow_blank=True)
    ws.add_data_validation(status)
    fcol = get_column_letter(columns.index("Final Status") + 1)
    status.add(f"{fcol}2:{fcol}{last}")
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = f"A1:{get_column_letter(len(columns))}{last}"
    _autosize(ws, columns)


def _pco_bom_sheet(ws, run: Run, review: ReviewBundle | None = None) -> None:
    root = run.metadata.input_root
    _write_header(ws, 1, PCO_BOM_COLUMNS)
    rows = [r for r in run.results if r.check.value == "PCO_BOM"]
    for i, r in enumerate(rows, start=2):
        a, b = r.source_a, r.source_b
        att = a.attributes if a else {}
        is_code = att.get("kind") == "affected_code"
        types = "; ".join(d.type.value for d in r.discrepancies)
        details = "; ".join(d.detail for d in r.discrepancies)
        actions = "; ".join(dict.fromkeys(d.recommended_action for d in r.discrepancies if d.recommended_action))
        sev = r.severity.value if r.severity else ""
        pco_number = Path(a.evidence.file).name if a else ""
        values = [
            r.row_id, r.sku, r.check.value, pco_number,
            "COVERAGE" if is_code else att.get("change_key", ""), "AFFECTED CODE" if is_code else att.get("change_kind", ""),
            att.get("item_actual", "") if not is_code else "", att.get("item_proposed", "") if not is_code else (a.item_number if a else ""), a.description if a else "",
            att.get("qty_proposed", ""), att.get("seq_proposed", ""), a.evidence.locator if a else "",
            _display_path(b.evidence.file, root) if b else "", b.evidence.page if b else None, b.evidence.locator if b else "", b.item_number if b else "", b.description if b else "", _qty(b), b.oper_seq if b else "",
            r.classification.value, types, details, sev, actions, "Y" if r.requires_validation else "N", r.explanation,
            *_review_cells(review, r),
        ]
        for c, v in enumerate(values, start=1):
            ws.cell(row=i, column=c, value=v).border = BORDER
        ws.cell(row=i, column=PCO_BOM_COLUMNS.index("Classification") + 1).fill = CLASS_FILLS[r.classification.value]
        if sev:
            ws.cell(row=i, column=PCO_BOM_COLUMNS.index("Severity") + 1).fill = SEVERITY_FILLS[sev]
        for name in ("Explanation", "Discrepancy Detail", "Recommended Action"):
            ws.cell(row=i, column=PCO_BOM_COLUMNS.index(name) + 1).alignment = WRAP
    _finish_results_sheet(ws, PCO_BOM_COLUMNS, len(rows))


def _coverage_sheet(ws, run: Run) -> None:
    _write_header(ws, 1, COVERAGE_COLUMNS)
    for i, c in enumerate(run.coverage, start=2):
        for col, v in enumerate([c.kind, c.sku, c.status, c.detail, c.source], start=1):
            cell = ws.cell(row=i, column=col, value=v)
            cell.border = BORDER
        if c.status != "OK":
            ws.cell(row=i, column=3).fill = SEVERITY_FILLS["BLOCKER" if "BOM" in c.status else "MAJOR"]
    ws.freeze_panes = "A2"
    ws.column_dimensions["A"].width = 20
    ws.column_dimensions["B"].width = 20
    ws.column_dimensions["C"].width = 16
    ws.column_dimensions["D"].width = 100
    ws.column_dimensions["E"].width = 50


def _bom_label_sheet(ws, run: Run) -> None:
    _pairing_sheet(ws, run, "BOM_LABEL")


def _pairing_sheet(ws, run: Run, check_type: str, review: ReviewBundle | None = None) -> None:
    root = run.metadata.input_root
    _write_header(ws, 1, BOM_LABEL_COLUMNS)
    rows = [r for r in run.results if r.check.value == check_type]
    for i, r in enumerate(rows, start=2):
        a, b = r.source_a, r.source_b
        types = "; ".join(d.type.value for d in r.discrepancies)
        details = "; ".join(d.detail for d in r.discrepancies)
        actions = "; ".join(dict.fromkeys(d.recommended_action for d in r.discrepancies if d.recommended_action))
        sev = r.severity.value if r.severity else ""
        values = [
            r.row_id, r.sku, r.check.value,
            _display_path(a.evidence.file, root) if a else "", a.evidence.page if a else None, a.evidence.locator if a else "", a.item_number if a else "", a.description if a else "", _qty(a), a.category.value if a else "",
            _display_path(b.evidence.file, root) if b else "", b.evidence.page if b else None, b.evidence.locator if b else "", b.item_number if b else "", b.description if b else "", _qty(b),
            r.normalized_a, r.normalized_b, r.classification.value, r.match_level.value, r.score, r.relationship_id,
            types, details, sev, actions, "Y" if r.requires_validation else "N", r.explanation,
            *_review_cells(review, r),
        ]
        for c, v in enumerate(values, start=1):
            cell = ws.cell(row=i, column=c, value=v)
            cell.border = BORDER
        ws.cell(row=i, column=BOM_LABEL_COLUMNS.index("Classification") + 1).fill = CLASS_FILLS[r.classification.value]
        if sev:
            ws.cell(row=i, column=BOM_LABEL_COLUMNS.index("Severity") + 1).fill = SEVERITY_FILLS[sev]
        for name in ("Explanation", "Discrepancy Detail", "Recommended Action"):
            ws.cell(row=i, column=BOM_LABEL_COLUMNS.index(name) + 1).alignment = WRAP
    _finish_results_sheet(ws, BOM_LABEL_COLUMNS, len(rows))


def _documents_sheet(ws, run: Run) -> None:
    root = run.metadata.input_root
    headers = ["Document ID", "Type", "File", "SHA-256", "SKU / REF", "Parser", "Parser Version", "Items", "Warnings"]
    _write_header(ws, 1, headers)
    row = 1
    for d in run.documents:
        row += 1
        for c, v in enumerate([d.id, d.doc_type.value, _display_path(d.path, root), d.sha256, d.sku, d.parser_name, d.parser_version, len(d.items), "; ".join(d.warnings)], start=1):
            ws.cell(row=row, column=c, value=v).border = BORDER
    row += 2
    ws.cell(row=row, column=1, value="BOM lines excluded from the BOM ↔ Label comparison (and why)").font = SECTION_FONT
    row += 1
    ex_headers = ["SKU", "Item", "Description", "Quantity", "Category", "Reason", "Locator", "File"]
    _write_header(ws, row, ex_headers)
    for d in run.documents:
        if d.doc_type is not DocType.BOM:
            continue
        for item in d.items:
            ok, reason = is_comparable_bom_item(item)
            if ok:
                continue
            row += 1
            for c, v in enumerate([d.sku, item.item_number, item.description, _qty(item), item.category.value, reason, item.evidence.locator, _display_path(d.path, root)], start=1):
                ws.cell(row=row, column=c, value=v).border = BORDER
    ws.freeze_panes = "A2"
    _autosize(ws, headers, default=18)
    ws.column_dimensions["D"].width = 66
    ws.column_dimensions["F"].width = 60
    ws.column_dimensions["I"].width = 60


def _relationships_sheet(ws, run: Run) -> None:
    headers = ["ID", "Version", "Canonical", "Aliases", "Scope", "Doc Types", "Item Anchors", "Provenance", "Created By", "Created At", "Active", "Notes", "Used In This Run", "Rows Using It"]
    _write_header(ws, 1, headers)
    used = set(run.relationships_used)
    for i, r in enumerate(run.terminology_snapshot, start=2):
        values = [r["id"], r.get("version", 1), r["canonical"], "; ".join(r.get("aliases", [])), r.get("scope"), "; ".join(r.get("doc_types", [])), "; ".join(r.get("item_anchors", [])), r.get("provenance"), r.get("created_by"), r.get("created_at"), "Y" if r.get("active") else "N", r.get("notes", ""), "Y" if r["id"] in used else "N", run.relationship_usage.get(r["id"], 0)]
        for c, v in enumerate(values, start=1):
            ws.cell(row=i, column=c, value=v).border = BORDER
    ws.cell(row=len(run.terminology_snapshot) + 3, column=1, value=f"Snapshot version (SHA-256 of all relationships above): {run.metadata.terminology_version}")
    ws.freeze_panes = "A2"
    _autosize(ws, headers, default=18)
    ws.column_dimensions["D"].width = 50
    ws.column_dimensions["L"].width = 60


def _metadata_sheet(ws, run: Run) -> None:
    m = run.metadata
    rows: list[tuple[str, Any]] = [("Run ID", m.run_id), ("Timestamp", m.timestamp.isoformat(timespec="seconds")), ("Tool version", m.tool_version)]
    rows += [(f"Parser version: {k}", v) for k, v in m.parser_versions.items()]
    rows += [(f"Threshold: {k}", v) for k, v in m.thresholds.model_dump().items()]
    rows += [("Terminology version (SHA-256)", m.terminology_version), ("Terminology relationship count", m.terminology_count), ("Input root", m.input_root), ("Relationships used in this run", ", ".join(run.relationships_used) or "none")]
    ws.cell(row=1, column=1, value="Run metadata").font = TITLE_FONT
    for i, (k, v) in enumerate(rows, start=2):
        ws.cell(row=i, column=1, value=k).font = SECTION_FONT
        ws.cell(row=i, column=2, value=v)
    row = len(rows) + 3
    ws.cell(row=row, column=1, value="Input files").font = SECTION_FONT
    row += 1
    _write_header(ws, row, ["File", "Doc Type", "SHA-256", "Size (bytes)"])
    for f in m.inputs:
        row += 1
        for c, v in enumerate([f.path, f.doc_type or "unrecognised", f.sha256, f.size_bytes], start=1):
            ws.cell(row=row, column=c, value=v).border = BORDER
    row += 2
    ws.cell(row=row, column=1, value="Capabilities").font = SECTION_FONT
    row += 1
    _write_header(ws, row, ["Capability", "Status"])
    for k, v in m.capabilities.items():
        row += 1
        ws.cell(row=row, column=1, value=k).border = BORDER
        ws.cell(row=row, column=2, value=v).border = BORDER
    ws.column_dimensions["A"].width = 60
    ws.column_dimensions["B"].width = 70
    ws.column_dimensions["C"].width = 66


def _audit_sheet(ws, run: Run) -> None:
    headers = ["Timestamp", "Actor", "Action", "Detail"]
    _write_header(ws, 1, headers)
    for i, e in enumerate(run.audit_log, start=2):
        for c, v in enumerate([e.timestamp.isoformat(timespec="seconds"), e.actor, e.action, e.detail], start=1):
            ws.cell(row=i, column=c, value=v).border = BORDER
    ws.column_dimensions["A"].width = 26
    ws.column_dimensions["C"].width = 20
    ws.column_dimensions["D"].width = 110


def _accuracy_sheet(ws, metrics: dict[str, Any]) -> None:
    ws.cell(row=1, column=1, value="Measured accuracy against ground truth").font = TITLE_FONT
    row = 3
    for k, v in metrics.get("overall", {}).items():
        ws.cell(row=row, column=1, value=k).font = SECTION_FONT
        ws.cell(row=row, column=2, value=v)
        row += 1
    for key in ("classification_accuracy", "pairing_accuracy", "ref_check_accuracy", "false_missing", "scored_rows"):
        if key in metrics:
            ws.cell(row=row, column=1, value=key).font = SECTION_FONT
            ws.cell(row=row, column=2, value=metrics[key])
            row += 1
    row += 1
    ws.cell(row=row, column=1, value="Per discrepancy type").font = SECTION_FONT
    row += 1
    _write_header(ws, row, ["Type", "TP", "FP", "FN", "Precision", "Recall"])
    for t, c in metrics.get("per_type", {}).items():
        row += 1
        for col, v in enumerate([t, c.get("tp"), c.get("fp"), c.get("fn"), c.get("precision"), c.get("recall")], start=1):
            ws.cell(row=row, column=col, value=v).border = BORDER
    if metrics.get("per_sku"):
        row += 2
        ws.cell(row=row, column=1, value="Per SKU").font = SECTION_FONT
        row += 1
        _write_header(ws, row, ["SKU", "Expected rows", "Pairing correct", "Classification correct", "TP", "FP", "FN"])
        for sku, c in metrics["per_sku"].items():
            row += 1
            for col, v in enumerate([sku, c.get("expected_rows"), c.get("pairing_correct"), c.get("classification_correct"), c.get("tp"), c.get("fp"), c.get("fn")], start=1):
                ws.cell(row=row, column=col, value=v).border = BORDER
    if metrics.get("mismatches"):
        row += 2
        ws.cell(row=row, column=1, value="Disagreements with ground truth").font = SECTION_FONT
        for m in metrics["mismatches"]:
            row += 1
            ws.cell(row=row, column=1, value=m)
    ws.column_dimensions["A"].width = 40
    ws.column_dimensions["B"].width = 16
