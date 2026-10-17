"""Value-only comparison workbook, projected from the same snapshot as Quick view."""

import re
from datetime import datetime, timezone
from pathlib import Path

from openpyxl import Workbook
from openpyxl.cell.cell import ILLEGAL_CHARACTERS_RE
from openpyxl.styles import Alignment, Font, PatternFill
from openpyxl.utils import get_column_letter

from kaizen.reporting.workbook import save_workbook

KINDS = ("BOM", "LABEL", "DRAWING")
NAMES = {"BOM": "BOM", "LABEL": "Label", "DRAWING": "Diagram"}
STATUS = {"matched": "Matched", "attention": "Missing or needs review", "neutral": "Not compared"}
COLORS = {"matched": ("E8F4EA", "216335"), "attention": ("FFF0DA", "8A4B08"), "neutral": ("F3F5F7", "475569")}


def _append(sheet, values):
    sheet.append([ILLEGAL_CHARACTERS_RE.sub("", v) if isinstance(v, str) else v for v in values])


def export_quick_view(payloads: list[dict], path: Path, *, run_id: str, run_name: str) -> Path:
    """All entries, one comparison sheet per SKU; Sources retains every original item.

    No formulas, filesystem paths, document bytes or independent matching logic.
    Screen pagination/search never truncates this export.
    """
    book = Workbook()
    info = book.active
    info.title = "Read me"
    for row in [
        ["Quick view comparison", run_name or run_id], ["Run ID", run_id],
        ["Exported (UTC)", datetime.now(timezone.utc).isoformat()],
        ["Scope", "All entries for the selected SKUs, including entries not compared. Screen filters are not applied."],
        ["Status", "Current saved review snapshot. This workbook does not approve results and cannot be imported as decisions."],
        ["Green", "Accepted counterpart with no unresolved finding."],
        ["Orange", "Missing counterpart or review needed. Unmapped does not establish physical absence."],
        ["Grey", "Not compared; no accepted match is claimed."],
        ["Source detail", "Sources retains original extracted text, quantities, filenames, pages and locators. Long text continues in additional rows."],
    ]:
        _append(info, row)
    info.column_dimensions["A"].width = 24
    info.column_dimensions["B"].width = 100
    sources = book.create_sheet("Sources")
    sources.append(["SKU", "Comparison row", "Source", "Status", "Item code", "Description", "Quantity", "Unit", "Filename", "Page", "Locator", "Original text", "Text part"])
    for index, payload in enumerate(payloads, 1):
        sku = payload["sku"]
        title = re.sub(r"[\\/*?:\[\]]", "_", str(sku))
        sheet = book.create_sheet(f"{index} {title}"[:31])
        sheet.append([f"SKU {index}: {sku}", "BOM", "Label", "Diagram", "Saved checks / findings"])
        for row in payload["rows"]:
            values = [STATUS[row["status"]]]
            for kind in KINDS:
                cell = row["cells"][kind]
                if not cell:
                    available = any(d["doc_type"] == kind for d in payload["documents"])
                    values.append("Not compared" if row["status"] == "neutral" else "No counterpart in this row" if available else "Document unavailable")
                    continue
                parts = [STATUS[cell["status"]]]
                for item in cell["items"]:
                    quantity = f"Qty {item['quantity']} {item.get('uom') or ''}" if item.get("quantity") is not None else ""
                    location = f"{item['file_name']} · p{item['page']}" if item.get("page") else item["file_name"]
                    parts.append(" · ".join(str(v) for v in (item.get("item_number"), item["description"], quantity, location, item.get("locator")) if v))
                    raw = item.get("raw_text") or ""
                    # Excel has a 32,767-character cell limit. Preserve long values
                    # in the source ledger instead of silently losing their tail.
                    chunks = [raw[n:n + 32000] for n in range(0, len(raw), 32000)] or [""]
                    description = item["description"]
                    descriptions = [description[n:n + 32000] for n in range(0, len(description), 32000)] or [""]
                    for part in range(max(len(chunks), len(descriptions))):
                        quantity = float(item["quantity"]) if item.get("quantity") is not None and part == 0 else None
                        _append(sources, [sku, row["id"], NAMES[kind], STATUS[cell["status"]], item.get("item_number"), descriptions[part] if part < len(descriptions) else "", quantity, item.get("uom"), item["file_name"], item.get("page"), item.get("locator"), chunks[part] if part < len(chunks) else "", part + 1])
                content = "\n\n".join(parts)
                values.append(content if len(content) <= 32767 else content[:32000] + "\n… Full entries retained on Sources.")
            checks = "\n\n".join(f"{c['check']} · {c['state']} · {c['explanation']}" + "".join(f"\n{f['type']}: {f['detail']}" for f in c["findings"]) for c in row["checks"])
            values.append(checks[:32767])
            _append(sheet, values)
            for col, status in [(1, row["status"]), *[(n + 2, row["cells"][k]["status"] if row["cells"][k] else row["status"]) for n, k in enumerate(KINDS)]]:
                background, foreground = COLORS[status]
                sheet.cell(sheet.max_row, col).fill = PatternFill("solid", fgColor=background)
                sheet.cell(sheet.max_row, col).font = Font(name="Calibri", size=11, color=foreground)
        for col, width in enumerate([26, 60, 60, 60, 70], 1):
            sheet.column_dimensions[get_column_letter(col)].width = width
    for col, width in enumerate([24, 35, 14, 26, 24, 60, 14, 12, 32, 10, 35, 100, 12], 1):
        sources.column_dimensions[get_column_letter(col)].width = width
    for sheet in book:
        sheet.freeze_panes = "B2" if sheet not in (info, sources) else "A2"
        sheet.auto_filter.ref = sheet.dimensions
        for cells in sheet:
            for cell in cells:
                if isinstance(cell.value, str):
                    cell.value = ILLEGAL_CHARACTERS_RE.sub("", cell.value)
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if cell.row == 1:
                    cell.fill = PatternFill("solid", fgColor="23364D")
                    cell.font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
                elif not cell.has_style or cell.font.name is None:
                    cell.font = Font(name="Calibri", size=11)
        sheet.sheet_view.zoomScale = 80
    return save_workbook(book, path)
