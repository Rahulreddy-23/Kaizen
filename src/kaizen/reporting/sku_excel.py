"""Reviewer downloads: a compact summary or complete comparisons, one worksheet per SKU."""

import json
import math
import re
from collections import defaultdict
from typing import TYPE_CHECKING, Literal

from openpyxl import Workbook
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from kaizen.models import CheckResult, Run
from kaizen.review.progress import SEVERITY_RANK

if TYPE_CHECKING:
    from kaizen.reporting.excel import ReviewBundle

CHECK_LABELS = {"BOM_LABEL": "BOM to label", "BOM_DRAWING": "BOM to drawing", "LABEL_DRAWING": "Label to drawing", "PCO_BOM": "PCO to BOM", "LABEL_REVISION": "Label revision", "COVERAGE": "Document coverage"}
SUMMARY_COLUMNS = ["Check", "Component", "Current Classification", "Review State", "Finding / Next Step", "Reviewer Decision", "Reviewer Comment", "Classification", "Row ID", "Reviewer 2 Decision", "Reviewer 2 Comment"]
COMPLETE_FRONT = ["Check", "Source A Item", "Source A Description", "Source B Description", "Source A Quantity", "Source B Quantity", "Current Classification",
                  "Review State", "Reviewer Decision", "Reviewer Comment"]
COLUMN_LABELS = {"Current Classification": "Result", "Review State": "Status", "Finding / Next Step": "Finding", "Source A Item": "Item",
                 "Source A Description": "Description A", "Source B Description": "Description B", "Source A Quantity": "Qty A", "Source B Quantity": "Qty B"}
STATE_LABELS = {"ENGINE_RECOMMENDED": "Not reviewed", "REVIEWED": "Awaiting approval", "FINALIZED": "Approved", "REVIEWER_1_COMPLETE": "Reviewer 1 complete", "REVIEWER_2_COMPLETE": "Reviewer 2 complete", "AGREED": "Agreed", "DISAGREEMENT": "Disagreement"}
EXTRA_COLUMNS = ["Current Classification", "Pending Review", "Open Finding", "Role", "Source A UOM", "Source B UOM", "Source A Oper Seq", "Source B Oper Seq",
                 "Source A Confidence", "Source B Confidence", "Source A Raw Text", "Source B Raw Text", "Source A Attributes", "Source B Attributes",
                 "Candidate Score", "Source A BBox", "Source B BBox", "Source A Extraction Method", "Source B Extraction Method", "Approved By", "Approved At", "Approval Note", "Self-approved"]
RESERVED = {"run_metadata", "action_items", "coverage", "manual_checklist", "documents", "relationships_used", "audit_log", "accuracy", "approval history"}
WIDTHS = {"Check": 18, "Component": 38, "Classification": 16, "Current Classification": 15, "Review State": 19, "Finding / Next Step": 40,
          "Reviewer Decision": 24, "Reviewer Comment": 30, "Reviewer 2 Decision": 24, "Reviewer 2 Comment": 30,
          "Source A Description": 36, "Source B Description": 36, "Source A File": 26, "Source B File": 26, "Source A Quantity": 10, "Source B Quantity": 10,
          "Explanation": 65, "Discrepancy Detail": 60, "Recommended Action": 50, "Source A Raw Text": 45, "Source B Raw Text": 45,
          "Source A Attributes": 45, "Source B Attributes": 45, "Approval Note": 40}


def _sheet_name(sku: str, used: set[str]) -> str:
    name = re.sub(r"[\[\]:*?/\\\x00-\x1f]", "-", sku).strip().strip("'")[:31] or "Unknown SKU"
    candidate, number = name, 2
    while candidate.casefold() in used:
        suffix = f" ({number})"
        candidate = name[:31 - len(suffix)] + suffix
        number += 1
    used.add(candidate.casefold())
    return candidate


def _current(row: CheckResult, review: "ReviewBundle | None") -> dict:
    if review and row.row_id in review.current_rows:
        return review.current_rows[row.row_id]
    discrepancies = [d.model_dump(mode="json") for d in row.discrepancies if d.severity.value != "INFO" or row.requires_validation]
    return {"effective_classification": row.classification.value, "state": "ENGINE_RECOMMENDED", "pending_review": row.requires_validation,
            "unresolved": bool(discrepancies), "current_discrepancies": discrepancies, "current_severity": row.severity.value if discrepancies and row.severity else None, "final": None}


def _header(ws, row, columns):
    for col, label in enumerate(columns, 1):
        cell = ws.cell(row, col, COLUMN_LABELS.get(label, label))
        cell.fill = PatternFill("solid", fgColor="F2F2F2")
        cell.font = Font(name="Calibri", size=11, bold=True, color="000000")
        cell.border = Border(bottom=Side(style="thin", color="BFBFBF"))
        cell.alignment = Alignment(horizontal="left", vertical="center", wrap_text=True)
    ws.row_dimensions[row].height = 28


def _finding(row, current):
    parts = [d["detail"] for d in current["current_discrepancies"]]
    if not parts and row.role in ("header", "reference", "coverage"):
        parts.append(row.explanation)
    return "\n".join(dict.fromkeys(parts))


def _values(row, run, review, current, detail):
    from kaizen.reporting.excel import BOM_LABEL_COLUMNS, _display_path, _qty, _review_cells

    a, b = row.source_a, row.source_b
    review_values = dict(zip(BOM_LABEL_COLUMNS[-11:], _review_cells(review, row)))
    if detail == "summary":
        texts = [f"{item.item_number + ': ' if item.item_number else ''}{item.description}" for item in (a, b) if item]
        component = "\n".join(dict.fromkeys(texts)) or row.explanation
        return {"Row ID": row.row_id, "Check": CHECK_LABELS.get(row.check.value, row.check.value), "Component": component,
                "Classification": row.classification.value, "Current Classification": current["effective_classification"],
                "Finding / Next Step": _finding(row, current), **review_values,
                "Review State": STATE_LABELS.get(current["state"], current["state"])}
    values = {"Row ID": row.row_id, "SKU": row.sku, "Check": CHECK_LABELS.get(row.check.value, row.check.value), "Role": row.role,
              "Normalized A": row.normalized_a, "Normalized B": row.normalized_b, "Classification": row.classification.value, "Match Level": row.match_level.value,
              "Score": row.score, "Candidate Score": row.candidate_score, "Relationship ID": row.relationship_id, "Discrepancy Type": "; ".join(d.type.value for d in row.discrepancies),
              "Discrepancy Detail": "; ".join(d.detail for d in row.discrepancies), "Severity": row.severity.value if row.severity else "",
              "Recommended Action": "; ".join(dict.fromkeys(d.recommended_action for d in row.discrepancies if d.recommended_action)),
              "Requires Validation": "Y" if row.requires_validation else "N", "Explanation": row.explanation,
              "Current Classification": current["effective_classification"], "Pending Review": "Y" if current["pending_review"] else "N", "Open Finding": "Y" if current["unresolved"] else "N",
              **review_values, "Review State": STATE_LABELS.get(current["state"], current["state"])}
    for side, item in (("A", a), ("B", b)):
        values.update({f"Source {side} File": _display_path(item.evidence.file, run.metadata.input_root) if item else "",
                       f"Source {side} Page": item.evidence.page if item else None, f"Source {side} Locator": item.evidence.locator if item else "",
                       f"Source {side} Item": item.item_number if item else "", f"Source {side} Description": item.description if item else "",
                       f"Source {side} Quantity": _qty(item), f"Source {side} Category": item.category.value if item else "",
                       f"Source {side} UOM": item.uom if item else "", f"Source {side} Oper Seq": item.oper_seq if item else "",
                       f"Source {side} Confidence": item.extraction_confidence if item else None, f"Source {side} Raw Text": item.evidence.raw_text if item else "",
                       f"Source {side} Attributes": json.dumps(item.attributes, ensure_ascii=False, default=str) if item else "",
                       f"Source {side} BBox": json.dumps(item.evidence.bbox.model_dump()) if item and item.evidence.bbox else "",
                       f"Source {side} Extraction Method": item.evidence.extraction_method if item else ""})
    final = review.finals.get(row.row_id) if review else None
    values.update({"Approved By": final.finalized_by if final else "", "Approved At": final.finalized_at if final else "",
                   "Approval Note": final.note if final else "", "Self-approved": "Yes" if final and final.note.startswith("Self-approved.") else ""})
    return values


def _overview(ws, run, sku, rows, current, detail):
    if detail == "complete":
        return 1
    headers = ["Check", "Rows", "Pending", "Findings", "Approved"]
    _header(ws, 1, headers)
    grouped = defaultdict(list)
    for row in rows:
        if row.role != "exempt":
            grouped[row.check.value].append(row)

    def counts(items):
        return [len(items), sum(current[r.row_id]["pending_review"] for r in items), sum(current[r.row_id]["unresolved"] for r in items),
                sum(bool(current[r.row_id]["final"]) for r in items)]

    line = 1
    for check, items in grouped.items():
        line += 1
        for col, value in enumerate([CHECK_LABELS.get(check, check), *counts(items)], 1):
            ws.cell(line, col, value)
    line += 1
    for col, value in enumerate(["Total", *counts([r for r in rows if r.role != "exempt"])], 1):
        ws.cell(line, col, value).font = Font(name="Calibri", size=11, bold=True)
    return line + 2


def sku_workbook(run: Run, detail: Literal["summary", "complete"], review: "ReviewBundle | None" = None) -> Workbook:
    from kaizen.reporting.excel import BOM_LABEL_COLUMNS, DECISIONS

    if detail not in ("summary", "complete"):
        raise ValueError("Excel detail must be summary or complete")
    wb = Workbook()
    wb.remove(wb.active)
    by_sku = {g.sku: [] for g in run.groups}
    for row in run.results:
        by_sku.setdefault(row.sku, []).append(row)
    if not by_sku:
        by_sku["No SKUs"] = []
    used = set(RESERVED)
    columns = SUMMARY_COLUMNS if detail == "summary" else list(dict.fromkeys([*COMPLETE_FRONT, *BOM_LABEL_COLUMNS, "Source B Category", *EXTRA_COLUMNS]))
    second_reviewer = bool(review and any(2 in slots for slots in review.decisions.values()))
    for sku, rows in by_sku.items():
        ws = wb.create_sheet(_sheet_name(sku, used))
        ws.sheet_view.showGridLines = True
        ws.sheet_properties.outlinePr.summaryBelow = False
        current = {r.row_id: _current(r, review) for r in rows}
        header = _overview(ws, run, sku, rows, current, detail)
        _header(ws, header, columns)
        selected = [r for r in rows if detail == "complete" or (r.role != "exempt" and (current[r.row_id]["pending_review"] or current[r.row_id]["unresolved"]))]
        if detail == "summary":
            selected.sort(key=lambda r: (SEVERITY_RANK.get(current[r.row_id].get("current_severity"), 4), r.check.value, r.row_id))
        for col, label in enumerate(columns, 1):
            ws.column_dimensions[get_column_letter(col)].width = WIDTHS.get(label, max(16, min(len(label) + 2, 30)))
            if (detail == "summary" and label in ("Row ID", "Classification")) or (label.startswith("Reviewer 2 ") and not second_reviewer):
                ws.column_dimensions[get_column_letter(col)].hidden = True
                if detail == "summary" and label == "Classification":
                    ws.column_dimensions[get_column_letter(col)].outlineLevel = 1
            if detail == "complete" and col > len(COMPLETE_FRONT):
                ws.column_dimensions[get_column_letter(col)].outlineLevel = 1
                ws.column_dimensions[get_column_letter(col)].hidden = True
        if detail == "complete":
            ws.sheet_properties.outlinePr.summaryRight = False
            ws.column_dimensions[get_column_letter(len(COMPLETE_FRONT))].collapsed = True
        for index, row in enumerate(selected, header + 1):
            values = _values(row, run, review, current[row.row_id], detail)
            max_lines = 1
            for col, label in enumerate(columns, 1):
                value = values.get(label)
                cell = ws.cell(index, col, value)
                cell.font = Font(name="Calibri", size=11, color="000000")
                cell.alignment = Alignment(vertical="top", wrap_text=True)
                if isinstance(value, (int, float)):
                    cell.number_format = "0.0%" if label in ("Score", "Candidate Score", "Source A Confidence", "Source B Confidence") else "#,##0" if float(value).is_integer() else "#,##0.####"
                if value is not None and not ws.column_dimensions[get_column_letter(col)].hidden:
                    width = ws.column_dimensions[get_column_letter(col)].width - 2
                    max_lines = max(max_lines, sum(max(1, math.ceil(len(part) / width)) for part in str(value).split("\n")))
            ws.row_dimensions[index].height = min(409, max(18, max_lines * 14 + 2))
        last = header + len(selected)
        if selected:
            ws.auto_filter.ref = f"A{header}:{get_column_letter(len(columns))}{last}"
            for label in ("Reviewer Decision", "Reviewer 2 Decision"):
                column = get_column_letter(columns.index(label) + 1)
                validation = DataValidation(type="list", formula1=f'"{DECISIONS}"', allow_blank=True)
                validation.errorTitle, validation.error = "Choose a decision", "Use one of the listed review decisions."
                validation.showErrorMessage = True
                ws.add_data_validation(validation)
                validation.add(f"{column}{header + 1}:{column}{last}")
        else:
            ws.cell(header + 1, 1, "No rows need attention." if any(r.role != "exempt" for r in rows) and detail == "summary" else "No comparisons available for this SKU.")
        ws.freeze_panes = f"C{header + 1}"
        ws.sheet_view.zoomScale = 100
        for cells in ws:
            for cell in cells:
                if cell.value is not None:
                    cell.font = Font(name="Calibri", size=11, bold=cell.font.bold, color="000000")
        ws.sheet_properties.pageSetUpPr.fitToPage = detail == "summary"
        ws.page_setup.orientation, ws.page_setup.paperSize = "landscape", ws.PAPERSIZE_A3
        ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = (1 if detail == "summary" else 0), 0
        ws.print_title_rows = f"{header}:{header}"
        ws.print_area = f"A1:{get_column_letter(len(columns))}{max(last, header + 1)}"
    return wb


def plain_workbook(wb: Workbook) -> None:
    """Give supporting audit sheets the same ordinary table formatting."""
    for ws in wb:
        ws.sheet_view.showGridLines = True
        ws.sheet_view.zoomScale = 100
        ws.sheet_properties.tabColor = None
        for cells in ws:
            for cell in cells:
                if cell.value is None:
                    continue
                is_header = cell.fill.fgColor.type == "rgb" and cell.fill.fgColor.rgb[-6:] in ("1F3864", "F2F2F2")
                cell.fill = PatternFill("solid", fgColor="F2F2F2") if is_header else PatternFill()
                cell.border = Border(bottom=Side(style="thin", color="BFBFBF")) if is_header else Border()
                cell.font = Font(name="Calibri", size=11, bold=is_header or cell.font.bold, color="000000")
